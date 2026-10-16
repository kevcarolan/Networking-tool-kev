"""HTTP API for the master circuit list."""

import io
import json
from collections import defaultdict

from fastapi import APIRouter, Body, Depends, HTTPException, Query, Request, Response
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import JSONResponse
from pydantic import BaseModel
from pydantic import Field as PField
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.core import exports
from app.core.audit import audit
from app.core.auth import User, require, require_admin
from app.core.db import get_db, session_scope, utcnow
from app.core.models import Device
from app.tools.circuits.importer import (_HEADERS, FIELDS, LABELS, CircuitImportError, Mapping,
                                         NeedsMapping, ParsedRow, ParsedSheet, _norm_header,
                                         _validate, diff, guess, header_signature,
                                         match_switches, normalize_mac, normalize_port,
                                         normalize_switch, parse_csv, parse_workbook,
                                         read_head, row_key)
from app.tools.circuits import live
from app.tools.circuits.models import (ACTIVE, PENDING, SUPERSEDED, Circuit, CircuitField,
                                       CircuitImport, CircuitScan, ColumnMap, DeviceScanState,
                                       LiveArp, LiveMac, LiveNeighbor)
from app.tools.circuits.store import LiveIndex

router = APIRouter(prefix="/api/circuits", tags=["circuits"])

CIRCUITS_VIEW = require("circuits", "view")
CIRCUITS_MANAGE = require("circuits", "manage")

CSV_HEADERS = LABELS


def _custom(db: Session) -> list[CircuitField]:
    """The columns admins added, in order."""
    return list(db.scalars(select(CircuitField).order_by(CircuitField.position, CircuitField.id)))


def _custom_headers(db: Session) -> dict[str, str]:
    """Normalised column name -> x_<key>, for reading uploads."""
    return {_norm_header(f.label): f"x_{f.key}" for f in _custom(db)}


def _active(db: Session) -> CircuitImport | None:
    return db.scalar(select(CircuitImport).where(CircuitImport.status == ACTIVE))


def _rows(db: Session, import_id: int) -> list[Circuit]:
    return db.scalars(select(Circuit).where(Circuit.import_id == import_id)
                      .order_by(Circuit.row_no)).all()


def _matches(db: Session, rows) -> dict[str, tuple[str, list[int]]]:
    names = dict(db.execute(select(Device.id, Device.name)).all())
    return match_switches(names, {r.switch_norm for r in rows if r.switch_norm})


class SiteIndex:
    """What a circuit's site is: the list's Site column, else the site of the switch it is
    documented on, else of the switch it was last found on. Matched to the site list by
    name or code (``sites.py``)."""

    def __init__(self, db: Session):
        from app.tools.inventory.models import Asset, Site
        from app.tools.inventory.sites import _index

        self.sites = _index(list(db.scalars(select(Site))))
        self.device_site = {d.id: d.site for d in db.scalars(select(Device))}
        for a in db.scalars(select(Asset).where(Asset.device_id.is_not(None), Asset.site != "")):
            self.device_site[a.device_id] = a.site  # typed on the inventory row: it wins

    def resolve(self, listed: str, switch_id: int | None, live_id: int | None) -> dict:
        text, source = listed.strip(), "list"
        if not text and switch_id:
            text, source = (self.device_site.get(switch_id) or "").strip(), "switch"
        if not text and live_id:
            text, source = (self.device_site.get(live_id) or "").strip(), "network"
        site = self.sites.get(text.lower()) if text else None
        return {"site_text": text, "site_from": source if text else "",
                "site_id": site.id if site else None, "site_name": site.name if site else ""}


def _row_out(c: Circuit, matches: dict, names: dict[int, str], idx: LiveIndex | None = None,
             sites: SiteIndex | None = None) -> dict:
    status, ids = matches.get(c.switch_norm, ("none", []))
    out = {f: getattr(c, f) for f in FIELDS}
    out.update({f"x_{k}": v for k, v in json.loads(c.custom or "{}").items()})
    doc_id = ids[0] if status == "matched" else None
    out.update(id=c.id, row_no=c.row_no, port_norm=c.port_norm, warnings=c.warnings,
               match=status if c.switch_norm else "none",
               switch_device_id=doc_id, switch_device_name=names.get(doc_id) if doc_id else None)
    if idx is not None:
        out.update(_live(c, doc_id, names, idx))
    if sites is not None:
        out.update(sites.resolve(c.site, doc_id, out.get("live_device_id")))
    return out


def _live(c: Circuit, doc_id: int | None, names: dict[int, str], idx: LiveIndex) -> dict:
    """The circuit's live location from the last MAC/ARP tables (live.resolve)."""
    if not idx.has_data:
        return {"live_status": "", "live_status_label": "Not scanned yet", "live_device_id": None,
                "live_switch": "", "live_port": "", "live_vlan": "", "live_mac": c.mac_norm,
                "live_mac_source": "", "live_scanned_at": None, "last_seen_at": None,
                "last_seen_where": ""}
    r = live.resolve(c.mac_norm, c.ip.strip(), (doc_id, c.port_norm if c.port else ""),
                     idx.edges, idx.arp)
    dev = r["live_device_id"]
    st = idx.scanned.get(dev) if dev else None
    last = idx.last_seen.get(r["live_mac"]) if r["live_mac"] else None
    return {**r, "live_status_label": live.STATUS_LABELS.get(r["live_status"], ""),
            "live_switch": names.get(dev, "") if dev else "",
            "live_scanned_at": st.last_success if st else None,
            "last_seen_at": last.seen_at if last else None,
            "last_seen_where": f"{names.get(last.device_id, '?')} {last.port}" if last else ""}


def _names(db: Session) -> dict[int, str]:
    return dict(db.execute(select(Device.id, Device.name)).all())


def _scan_out(db: Session) -> dict:
    last = db.scalar(select(CircuitScan).order_by(CircuitScan.id.desc()).limit(1))
    full = db.scalar(select(CircuitScan).where(CircuitScan.trigger.not_like("upgrade-job:%"))
                     .where(CircuitScan.finished_at.is_not(None))
                     .order_by(CircuitScan.id.desc()).limit(1))
    return {"last": _scan_row(last) if last else None,
            "last_full": _scan_row(full) if full else None}


def _scan_row(s: CircuitScan) -> dict:
    return {"id": s.id, "trigger": s.trigger, "started_at": s.started_at,
            "finished_at": s.finished_at, "devices_ok": s.devices_ok,
            "devices_failed": s.devices_failed, "errors": s.errors}


def _import_out(imp: CircuitImport) -> dict:
    return {"id": imp.id, "filename": imp.filename, "sheet": imp.sheet, "status": imp.status,
            "uploaded_by": imp.uploaded_by, "uploaded_at": imp.uploaded_at,
            "committed_by": imp.committed_by, "committed_at": imp.committed_at,
            "row_count": imp.row_count, "warning_count": imp.warning_count,
            "summary": json.loads(imp.summary or "{}")}


# --- the current list -------------------------------------------------------

@router.get("")
def list_circuits(db: Session = Depends(get_db), _: User = Depends(CIRCUITS_VIEW)):
    imp = _active(db)
    scan = _scan_out(db)
    if imp is None:
        return {"import": None, "circuits": [], "scan": scan}
    rows = _rows(db, imp.id)
    matches = _matches(db, rows)
    names, idx = _names(db), LiveIndex(db)
    return {"import": _import_out(imp), "scan": scan,
            "circuits": [_row_out(c, matches, names, idx, SiteIndex(db)) for c in rows]}


@router.get("/export.csv")
def export_csv(db: Session = Depends(get_db), _: User = Depends(CIRCUITS_VIEW)):
    custom = [f"x_{f.key}" for f in _custom(db)]
    out = io.StringIO()
    w = exports.csv_writer(out)
    w.writerow([CSV_HEADERS[f] for f in FIELDS] + [f.label for f in _custom(db)]
               + ["Matched inventory device", "Linked site", "Warnings"] + LIVE_HEADERS)
    for r in _export_rows(db):
        w.writerow([r[f] for f in FIELDS] + [r.get(k, "") for k in custom]
                   + [r["switch_device_name"] or r["match"], r["site_name"], r["warnings"]]
                   + _live_cells(r))
    return Response(out.getvalue(), media_type="text/csv", headers={
        "Content-Disposition": 'attachment; filename="circuit-list.csv"'})


LIVE_HEADERS = ["Live status", "Live switch", "Live port", "Live VLAN", "MAC used",
                "Live data from", "Last seen", "Last seen on"]


def _fmt(t) -> str:
    return f"{t:%Y-%m-%d %H:%M}" if t else ""


def _live_cells(r: dict) -> list:
    mac = r["live_mac"] + (" (from ARP)" if r["live_mac_source"] == "arp" else "")
    return [r["live_status_label"], r["live_switch"], r["live_port"], r["live_vlan"], mac,
            _fmt(r["live_scanned_at"]), _fmt(r["last_seen_at"]), r["last_seen_where"]]


def _export_rows(db: Session) -> list[dict]:
    imp = _active(db)
    if imp is None:
        return []
    rows = _rows(db, imp.id)
    matches = _matches(db, rows)
    names, idx = _names(db), LiveIndex(db)
    sites = SiteIndex(db)
    return [_row_out(c, matches, names, idx, sites) for c in rows]


def site_counts(db: Session) -> dict[int, int]:
    """Circuits per site (site id -> count), for the Sites page."""
    counts: dict[int, int] = defaultdict(int)
    for r in _export_rows(db):
        if r["site_id"]:
            counts[r["site_id"]] += 1
    return counts


@router.get("/export-updated.xlsx")
def export_updated(db: Session = Depends(get_db), _: User = Depends(CIRCUITS_VIEW)):
    """The circuit list with DEVICE and PORT taken from the network where the circuit was
    found, plus the documented values and the live status. Uploading it again makes the
    live locations the documented ones."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Circuits"
    custom = _custom(db)
    ws.append([CSV_HEADERS[f] for f in FIELDS] + [f.label for f in custom]
              + ["Documented DEVICE", "Documented PORT"] + LIVE_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in _export_rows(db):
        found = r["live_status"] in (live.MATCHES, live.MOVED, live.FOUND)
        values = {**{f: r[f] for f in FIELDS}}
        if found:
            values["switch"], values["port"] = r["live_switch"], r["live_port"]
            values["vlan"] = r["vlan"] or r["live_vlan"]
        exports.append_row(ws, [values[f] for f in FIELDS]
                           + [r.get(f"x_{f.key}", "") for f in custom]
                           + [r["switch"], r["port"]] + _live_cells(r))
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), headers={
        "Content-Disposition": 'attachment; filename="circuit-list-updated.xlsx"'},
        media_type=exports.XLSX_TYPE)


# --- the columns: the built-in ones, the admins' own, and blank templates -----------------------

# Columns in exports that an upload ignores: a custom column can't take these names.
_EXPORT_ONLY = ["Matched inventory device", "Linked site", "Warnings", "Documented DEVICE", "Documented PORT",
                *LIVE_HEADERS]


def _fields(db: Session) -> list[dict]:
    return [{"key": f, "label": LABELS[f], "custom": False} for f in FIELDS] + [
        {"key": f"x_{c.key}", "label": c.label, "custom": True} for c in _custom(db)]


@router.get("/meta")
def meta(db: Session = Depends(get_db), user: User = Depends(CIRCUITS_VIEW)):
    return {"fields": _fields(db), "can_manage": user.can("circuits", "manage")}


class FieldIn(BaseModel):
    label: str = PField(min_length=1, max_length=100)


def _check_label(db: Session, label: str, exclude: int | None = None) -> str:
    label = label.strip()
    norm = _norm_header(label)
    if not norm:
        raise HTTPException(422, "Give the column a name with letters or digits")
    taken = set(_HEADERS) | {_norm_header(x) for x in _EXPORT_ONLY} | {
        _norm_header(c.label) for c in _custom(db) if c.id != exclude}
    if norm in taken:
        raise HTTPException(422, f"There is already a column called {label} (or the import "
                                 "reads that name as a built-in column)")
    return label


@router.post("/fields", status_code=201)
def add_field(body: FieldIn, db: Session = Depends(get_db), user: User = Depends(require_admin)):
    label = _check_label(db, body.label)
    base = key = "_".join(_norm_header(label).split())[:40]
    n = 2
    while db.scalar(select(CircuitField.id).where(CircuitField.key == key)):
        key, n = f"{base}_{n}", n + 1
    position = (db.scalar(select(func.max(CircuitField.position))) or 0) + 1
    db.add(CircuitField(key=key, label=label, position=position))
    db.flush()
    audit(db, user.username, "circuits.column_add", label)
    return {"key": f"x_{key}"}


def _field(db: Session, key: str) -> CircuitField:
    f = db.scalar(select(CircuitField).where(CircuitField.key == key.removeprefix("x_")))
    if f is None:
        raise HTTPException(404, "No such column")
    return f


@router.put("/fields/{key}")
def rename_field(key: str, body: FieldIn, db: Session = Depends(get_db),
                 user: User = Depends(require_admin)):
    f = _field(db, key)
    old, f.label = f.label, _check_label(db, body.label, exclude=f.id)
    audit(db, user.username, "circuits.column_change", f"{old} -> {f.label}")
    return {"key": f"x_{f.key}"}


@router.delete("/fields/{key}", status_code=204)
def delete_field(key: str, db: Session = Depends(get_db), user: User = Depends(require_admin)):
    f = _field(db, key)
    for c in db.scalars(select(Circuit).where(Circuit.custom.like(f'%"{f.key}"%'))):
        values = json.loads(c.custom or "{}")
        values.pop(f.key, None)
        c.custom = json.dumps(values)
    db.delete(f)
    audit(db, user.username, "circuits.column_delete", f.label)
    return Response(status_code=204)


TEMPLATE_NOTES = [
    "Fill in one row per circuit on the first sheet, under the column names, then upload the file "
    "on the Circuits page (Upload spreadsheet).",
    "Keep the column names as they are. Columns can be in any order, and you can delete the ones "
    "you don't use.",
    "Each row needs a Hardware Mac Address or an IP Address: the tools find the circuit on the "
    "network from it.",
    "DEVICE and PORT are the switch and port the circuit is documented on (e.g. ACC-SW01, Gi1/0/5). "
    "They can be left empty: the network scan fills in where each circuit is.",
    "VLAN: a number from 1 to 4094. MAC: any format (0011.2233.4455, 00:11:22:33:44:55, ...).",
    "The upload replaces the whole list: rows not in the file are removed from it. You see what "
    "changes before it is made live.",
]


@router.get("/template.{fmt}")
def template(fmt: str, columns: str = Query("", max_length=4000), db: Session = Depends(get_db),
             _: User = Depends(CIRCUITS_VIEW)):
    """A blank circuit list to fill in: the chosen columns (``columns``: keys separated by
    commas), or every column. The MAC and IP address columns are always in it."""
    if fmt not in ("csv", "xlsx"):
        raise HTTPException(404, "Use template.csv or template.xlsx")
    fields = {f["key"]: f["label"] for f in _fields(db)}
    keys = [k for k in dict.fromkeys(c.strip() for c in columns.split(",")) if k in fields] \
        or list(fields)
    for needed in ("ip", "mac"):
        if needed not in keys:
            keys.insert(0, needed)
    return exports.template(fmt, [fields[k] for k in keys], "circuit-list-template",
                            TEMPLATE_NOTES, sheet="Circuits")


@router.get("/for-device/{device_id}")
def for_device(device_id: int, db: Session = Depends(get_db), _: User = Depends(CIRCUITS_VIEW)):
    """Circuits on one switch, grouped by service and VLAN (for upgrade planning)."""
    device = db.get(Device, device_id)
    if device is None:
        raise HTTPException(404, "Device not found")
    return circuits_for_device(db, device)


@router.get("/ports/{device_id}")
def ports(device_id: int, db: Session = Depends(get_db), _: User = Depends(CIRCUITS_VIEW)):
    """One switch's ports: the MAC addresses its last scan saw on each (with the IP from any
    device's ARP table), the circuits the list puts there or that were found there, and a
    status per port. From the stored scan: no device is contacted."""
    import re

    from app.tools.circuits.importer import normalize_port

    device = db.get(Device, device_id)
    if device is None:
        raise HTTPException(404, "Device not found")
    arp = {}
    for a in db.scalars(select(LiveArp).order_by(LiveArp.seen_at)):
        arp[a.mac_norm] = a.ip  # the newest wins
    out: dict[str, dict] = {}

    def port(norm: str, label: str) -> dict:
        return out.setdefault(norm, {"port": label, "port_norm": norm, "macs": [], "uplink": False,
                                     "circuits": []})

    for m in db.scalars(select(LiveMac).where(LiveMac.device_id == device.id)):
        p = port(m.port_norm, m.port)
        p["uplink"] |= bool(m.uplink)
        p["macs"].append({"mac": m.mac_norm, "vlan": m.vlan, "ip": arp.get(m.mac_norm, "")})
    data = circuits_for_device(db, device)
    for group in data["groups"]:
        for c in group["circuits"]:
            item = {"id": c["id"], "device_name": c["device_name"], "service": c["service"],
                    "vlan": c["vlan"], "mac": c["mac"], "ip": c["ip"],
                    "live_status": c["live_status"], "live_status_label": c["live_status_label"],
                    "documented": f"{c['switch']} {c['port']}".strip(),
                    "found": f"{c['live_switch']} {c['live_port']}".strip()}
            if c["switch_device_id"] == device.id and c["port_norm"]:
                port(c["port_norm"], c["port"])["circuits"].append({**item, "here": "documented"})
            if c["live_device_id"] == device.id and c["live_port"]:
                norm = normalize_port(c["live_port"])
                if c["switch_device_id"] != device.id or norm != c["port_norm"]:
                    port(norm, c["live_port"])["circuits"].append({**item, "here": "found"})
    for p in out.values():
        documented = [c for c in p["circuits"] if c["here"] == "documented"]
        if p["uplink"]:
            p["status"] = "uplink"
        elif any(c["live_status"] in (live.MOVED, live.NOT_SEEN) for c in documented):
            p["status"] = "problem"  # its circuit isn't here any more
        elif any(c["here"] == "found" for c in p["circuits"]):
            p["status"] = "moved_here"
        elif documented:
            p["status"] = "ok"
        else:
            p["status"] = "unlisted" if p["macs"] else "empty"

    def order(p):  # natural order: gi1/0/2 before gi1/0/10
        return [int(x) if x.isdigit() else x for x in re.split(r"(\d+)", p["port_norm"])]

    state = db.get(DeviceScanState, device.id)
    return {"device": {"id": device.id, "name": device.name, "platform": device.platform},
            "scanned_at": state.last_success if state else None,
            "scan_error": state.last_error if state else "",
            "ports": sorted(out.values(), key=order)}


def circuits_for_device(db: Session, device: Device) -> dict:
    """Circuits the list puts on this device, or that were last seen on it.
    Also used by upgrade jobs to freeze the affected circuits."""
    imp = _active(db)
    if imp is None:
        return {"device": device.name, "import": None, "count": 0, "groups": []}
    rows = _rows(db, imp.id)
    matches = _matches(db, rows)
    names, idx, sites = _names(db), LiveIndex(db), SiteIndex(db)
    groups = defaultdict(list)
    count = 0
    for c in rows:
        r = _row_out(c, matches, names, idx, sites)
        if device.id not in (r["switch_device_id"], r["live_device_id"]):
            continue
        r["on_this_device"] = ("both" if r["switch_device_id"] == r["live_device_id"] else
                               "list" if r["switch_device_id"] == device.id else "live")
        count += 1
        groups[(c.service or "(no service)", c.vlan or r["live_vlan"] or "-")].append(r)
    return {"device": device.name, "import": _import_out(imp), "count": count,
            "groups": [{"service": s, "vlan": v, "circuits": items}
                       for (s, v), items in sorted(groups.items())]}


def circuits_seen(db: Session, device_id: int, edge_macs: list[list[str]]) -> list[dict]:
    """Circuits of the live list whose MAC is on one of this device's edge ports
    (edge_macs: [mac, port, port_norm, vlan]). Saved in an upgrade job's pre-check
    snapshot and compared after the reload."""
    imp = _active(db)
    if imp is None or not edge_macs:
        return []
    here = {m[0]: m for m in edge_macs}
    idx = LiveIndex(db)
    seen = []
    for c in _rows(db, imp.id):
        mac = c.mac_norm or (idx.arp.get(c.ip.strip(), "") if c.ip else "")
        if mac in here:
            _, port, port_norm, vlan = here[mac]
            seen.append({"name": c.device_name or c.eng_prefix or f"row {c.row_no}", "mac": mac,
                         "service": c.service, "vlan": vlan or c.vlan, "port": port,
                         "port_norm": port_norm})
    return seen


# --- found on the network: what is plugged in but not on the list ---------------------

def discovered(db: Session) -> list[dict]:
    """Every MAC the last scans saw on an edge port (not an uplink) that no row of the live
    list stands for (by its MAC, or its IP through ARP). One entry per MAC, with the IP from
    any ARP table and what the device says about itself over LLDP/CDP on that port."""
    imp = _active(db)
    listed: set[str] = set()
    if imp is not None:
        arp_by_ip = LiveIndex(db).arp
        for c in _rows(db, imp.id):
            if c.mac_norm:
                listed.add(c.mac_norm)
            elif c.ip.strip() and arp_by_ip.get(c.ip.strip()):
                listed.add(arp_by_ip[c.ip.strip()])
    ip_of: dict[str, str] = {}
    for a in db.scalars(select(LiveArp).order_by(LiveArp.seen_at)):
        ip_of[a.mac_norm] = a.ip  # the newest wins
    neighbour = {(n.device_id, n.local_norm): n
                 for n in db.scalars(select(LiveNeighbor).where(LiveNeighbor.infra.is_(False)))}
    names, sites = _names(db), SiteIndex(db)
    found: dict[str, dict] = {}
    for m in db.scalars(select(LiveMac).where(LiveMac.uplink.is_(False))
                        .order_by(LiveMac.device_id, LiveMac.port_norm)):
        if not m.mac_norm or m.mac_norm in listed:
            continue
        item = found.get(m.mac_norm)
        if item is not None:  # on more than one edge port (a switch with no LLDP below it)
            item["also_on"].append(f"{names.get(m.device_id, '?')} {m.port}")
            continue
        n = neighbour.get((m.device_id, m.port_norm))
        site = sites.resolve("", None, m.device_id)
        found[m.mac_norm] = {
            "mac": m.mac_norm, "ip": ip_of.get(m.mac_norm, ""), "vlan": m.vlan,
            "device_id": m.device_id, "switch": names.get(m.device_id, ""), "port": m.port,
            "port_norm": m.port_norm, "site_id": site["site_id"],
            "site": site["site_name"] or site["site_text"],
            "lldp_name": n.name if n else "", "lldp_platform": n.platform if n else "",
            "seen_at": m.seen_at, "also_on": []}
    return sorted(found.values(), key=lambda r: (r["site"], r["switch"], r["port_norm"], r["mac"]))


@router.get("/discovered")
def list_discovered(db: Session = Depends(get_db), _: User = Depends(CIRCUITS_VIEW)):
    return {"scan": _scan_out(db), "devices": discovered(db)}


DISCOVERED_HEADERS = ["Hardware Mac Address", "IP Address", "VLAN", "DEVICE", "PORT", "Site",
                      "Device Name", "Says it is (LLDP/CDP)", "Also seen on", "Last seen"]


@router.get("/discovered.xlsx")
def discovered_xlsx(db: Session = Depends(get_db), _: User = Depends(CIRCUITS_VIEW)):
    """What was found, in the list's own column names: fill in Device Name, Service and the
    rest in Excel and upload it (the extra columns are ignored)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Found on the network"
    ws.append(DISCOVERED_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in discovered(db):
        exports.append_row(ws, [r["mac"], r["ip"], r["vlan"], r["switch"], r["port"], r["site"],
                                r["lldp_name"], r["lldp_platform"], ", ".join(r["also_on"]),
                                _fmt(r["seen_at"])])
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), headers={
        "Content-Disposition": 'attachment; filename="found-on-the-network.xlsx"'},
        media_type=exports.XLSX_TYPE)


class AdoptIn(BaseModel):
    macs: list[str] = PField(min_length=1, max_length=5000)
    service: str = PField("", max_length=200)
    device_type: str = PField("", max_length=255)


@router.post("/discovered/adopt", status_code=201)
def adopt(body: AdoptIn, db: Session = Depends(get_db), user: User = Depends(CIRCUITS_MANAGE)):
    """Add found devices to the list: a pending import of the live list plus these rows,
    shown as a preview like an upload and made live with the same commit."""
    wanted = {normalize_mac(m) for m in body.macs} - {""}
    picked = [r for r in discovered(db) if r["mac"] in wanted]
    if not picked:
        raise HTTPException(422, "None of these devices is on the found list any more: scan again")
    imp = _active(db)
    rows: list[ParsedRow] = []
    for c in _rows(db, imp.id) if imp else []:
        rows.append(ParsedRow(row_no=c.row_no, values={f: getattr(c, f) for f in FIELDS},
                              switch_norm=c.switch_norm, port_norm=c.port_norm,
                              mac_norm=c.mac_norm,
                              warnings=[w for w in (c.warnings or "").split("; ") if w],
                              custom=json.loads(c.custom or "{}")))
    next_row = max((r.row_no for r in rows), default=1) + 1
    for i, r in enumerate(picked):
        values = {f: "" for f in FIELDS}
        values.update(mac=r["mac"], ip=r["ip"], vlan=r["vlan"], switch=r["switch"],
                      port=r["port"], device_name=r["lldp_name"], service=body.service.strip(),
                      device_type=body.device_type.strip())
        rows.append(ParsedRow(row_no=next_row + i, values=values,
                              switch_norm=normalize_switch(r["switch"]),
                              port_norm=normalize_port(r["port"]), mac_norm=r["mac"],
                              warnings=_validate(values)))
    labels = {f["key"]: f["label"] for f in _fields(db)}
    parsed = ParsedSheet(sheet="", header_row=0, columns=labels, rows=rows)
    db.commit()  # _store_pending opens its own session
    return _store_pending(parsed, f"Found on the network ({len(picked)} added)", user.username)


# --- reading the network (MAC and ARP tables) ---------------------------------

@router.get("/scans")
def list_scans(db: Session = Depends(get_db), _: User = Depends(CIRCUITS_VIEW)):
    rows = db.scalars(select(CircuitScan).order_by(CircuitScan.id.desc()).limit(50))
    return [_scan_row(s) for s in rows]


@router.post("/scan", status_code=202)
def scan_now(request: Request, device_id: int | None = Body(None, embed=True),
             db: Session = Depends(get_db), user: User = Depends(CIRCUITS_MANAGE)):
    """Read the MAC and ARP tables now (every device, or one) with the backup accounts."""
    svc = request.app.state.circuit_scan
    if device_id is not None and db.get(Device, device_id) is None:
        raise HTTPException(404, "Device not found")
    if not svc.scan_in_background(f"manual:{user.username}",
                                  [device_id] if device_id is not None else None):
        raise HTTPException(409, "A scan is already running")
    audit(db, user.username, "circuits.scan",
          f"device {device_id}" if device_id is not None else "all devices")
    return {"started": True}


# --- imports ----------------------------------------------------------------

@router.get("/imports")
def list_imports(db: Session = Depends(get_db), _: User = Depends(CIRCUITS_VIEW)):
    rows = db.scalars(select(CircuitImport).order_by(CircuitImport.id.desc()).limit(100))
    return [_import_out(i) for i in rows]


@router.get("/imports/{import_id}")
def get_import(import_id: int, db: Session = Depends(get_db), _: User = Depends(CIRCUITS_VIEW)):
    imp = db.get(CircuitImport, import_id) or _not_found()
    return _import_out(imp)


def _saved_maps(db: Session) -> dict[str, dict[int, str]]:
    return {m.signature: {int(k): v for k, v in json.loads(m.columns or "{}").items()}
            for m in db.scalars(select(ColumnMap))}


def _mapping(text: str) -> Mapping | None:
    """The ``map`` query parameter: {"sheet": "", "header_row": 0, "columns": {"3": "mac"}}."""
    if not text:
        return None
    try:
        raw = json.loads(text)
        cols = {int(k): str(v) for k, v in (raw.get("columns") or {}).items() if v and v != "ignore"}
        return Mapping(header_row=int(raw.get("header_row", 0)), columns=cols,
                       sheet=str(raw.get("sheet") or "")[:255])
    except (ValueError, TypeError, AttributeError):
        raise HTTPException(422, "The column choices could not be read") from None


@router.put("/imports/upload", status_code=201)
async def upload(request: Request, filename: str = Query(..., max_length=255),
                 map: str = Query("", max_length=20000), save: bool = False,  # noqa: A002
                 user: User = Depends(CIRCUITS_MANAGE)):
    """Upload a spreadsheet (raw .xlsx or .csv body). It is checked and kept as a pending
    import with a preview; nothing changes until it is committed.

    When no header row is recognised (and no saved layout matches), the answer is
    ``needs_mapping`` with the file's first rows: the person says which column is which and
    sends the file again with ``map`` (and ``save`` to remember the layout)."""
    limit = request.app.state.settings.circuits_max_upload_mb * 2**20
    body = bytearray()
    async for chunk in request.stream():
        body += chunk
        if len(body) > limit:
            raise HTTPException(413, f"File is larger than {limit // 2**20} MB")
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        parse = parse_workbook
    elif name.endswith(".csv"):
        parse = parse_csv
    else:
        raise HTTPException(422, "Upload an Excel workbook (.xlsx) or a CSV file")
    mapping = _mapping(map)
    with session_scope() as db:
        custom, saved, fields = _custom_headers(db), _saved_maps(db), _fields(db)
    try:
        parsed = await run_in_threadpool(parse, bytes(body), custom=custom, mapping=mapping,
                                         saved=saved)
    except NeedsMapping as exc:
        try:
            sheets = await run_in_threadpool(read_head, bytes(body), parse is parse_csv)
        except CircuitImportError as err:
            raise HTTPException(422, str(err)) from None
        for sh in sheets:
            sh["guess"] = guess(sh["rows"], custom)
        return JSONResponse({"needs_mapping": True, "message": str(exc), "sheets": sheets,
                             "fields": fields}, status_code=200)
    except CircuitImportError as exc:
        raise HTTPException(422, str(exc)) from None
    if mapping is not None and save:
        await run_in_threadpool(_save_map, parsed, filename, user.username)
    return await run_in_threadpool(_store_pending, parsed, filename, user.username)


def _save_map(parsed, filename: str, username: str) -> None:
    sig = header_signature(parsed.headers)
    with session_scope() as db:
        m = db.scalar(select(ColumnMap).where(ColumnMap.signature == sig)) or ColumnMap(signature=sig)
        m.name, m.created_by, m.created_at = filename[:255], username, utcnow()
        m.headers = json.dumps(parsed.headers)
        m.columns = json.dumps({str(k): v for k, v in parsed.mapping.items()})
        db.add(m)
        audit(db, username, "circuits.layout.save",
              f"{filename}: {len(parsed.mapping)} columns")


@router.get("/column-maps")
def column_maps(db: Session = Depends(get_db), _: User = Depends(CIRCUITS_VIEW)):
    """The spreadsheet layouts people matched on upload."""
    labels = {f["key"]: f["label"] for f in _fields(db)}
    out = []
    for m in db.scalars(select(ColumnMap).order_by(ColumnMap.created_at.desc())):
        headers = json.loads(m.headers or "[]")
        cols = {int(k): v for k, v in json.loads(m.columns or "{}").items()}
        out.append({"id": m.id, "name": m.name, "created_by": m.created_by,
                    "created_at": m.created_at,
                    "columns": [{"header": headers[c] if c < len(headers) else f"column {c + 1}",
                                 "field": f, "label": labels.get(f, f)}
                                for c, f in sorted(cols.items())]})
    return out


@router.delete("/column-maps/{map_id}", status_code=204)
def delete_column_map(map_id: int, db: Session = Depends(get_db),
                      user: User = Depends(CIRCUITS_MANAGE)):
    m = db.get(ColumnMap, map_id)
    if m is None:
        raise HTTPException(404, "Layout not found")
    db.delete(m)
    audit(db, user.username, "circuits.layout.delete", m.name)
    return Response(status_code=204)


def _store_pending(parsed, filename: str, username: str) -> dict:
    with session_scope() as db:
        # Only one pending import at a time: a new upload replaces the old preview.
        for old in db.scalars(select(CircuitImport).where(CircuitImport.status == PENDING)):
            db.delete(old)
        imp = CircuitImport(filename=filename[:255], sheet=parsed.sheet, uploaded_by=username,
                            row_count=len(parsed.rows),
                            warning_count=sum(bool(r.warnings) for r in parsed.rows))
        db.add(imp)
        db.flush()
        db.add_all(Circuit(import_id=imp.id, row_no=r.row_no, switch_norm=r.switch_norm,
                           port_norm=r.port_norm, mac_norm=r.mac_norm,
                           warnings="; ".join(r.warnings), custom=json.dumps(r.custom),
                           **r.values) for r in parsed.rows)
        db.flush()

        new_rows = _rows(db, imp.id)
        active = _active(db)
        old_rows = _rows(db, active.id) if active else []
        matches = _matches(db, new_rows)
        unmatched = sorted({k for k, (s, _) in matches.items() if s == "unmatched"})
        ambiguous = sorted({k for k, (s, _) in matches.items() if s == "ambiguous"})
        labels = {f["key"]: f["label"] for f in _fields(db)}
        compared = list(labels)

        def values(c):
            custom = json.loads(c.custom or "{}")
            return {**{f: getattr(c, f) for f in FIELDS},
                    **{k: custom.get(k[2:], "") for k in compared if k.startswith("x_")}}

        summary = {
            "header_row": parsed.header_row,
            "columns": parsed.columns,
            "labels": labels,
            "missing_columns": [label for k, label in labels.items() if k not in parsed.columns],
            "switches": len(matches),
            "unmatched_switches": unmatched,
            "ambiguous_switches": ambiguous,
            "rows_unmatched": sum(1 for c in new_rows
                                  if matches.get(c.switch_norm, ("none",))[0] != "matched"),
            "warnings_sample": [{"row": c.row_no, "warnings": c.warnings}
                                for c in new_rows if c.warnings][:100],
            "diff": diff([(row_key(values(c), c.mac_norm), values(c)) for c in old_rows],
                         [(row_key(values(c), c.mac_norm), values(c)) for c in new_rows],
                         fields=compared),
            "replaces_import": active.id if active else None,
        }
        imp.summary = json.dumps(summary, default=str)
        audit(db, username, "circuits.import.preview",
              f"{filename}: {imp.row_count} rows, {imp.warning_count} with warnings")
        return _import_out(imp)


@router.post("/imports/{import_id}/commit")
def commit(import_id: int, db: Session = Depends(get_db), user: User = Depends(CIRCUITS_MANAGE)):
    imp = db.get(CircuitImport, import_id) or _not_found()
    if imp.status != PENDING:
        raise HTTPException(409, "Only a pending import can be committed")
    for old in db.scalars(select(CircuitImport).where(CircuitImport.status == ACTIVE)):
        old.status = SUPERSEDED
    imp.status, imp.committed_by, imp.committed_at = ACTIVE, user.username, utcnow()
    d = json.loads(imp.summary or "{}").get("diff", {})
    audit(db, user.username, "circuits.import.commit",
          f"{imp.filename}: {imp.row_count} rows (+{d.get('added', 0)} "
          f"~{d.get('changed', 0)} -{d.get('removed', 0)})")
    return _import_out(imp)


@router.delete("/imports/{import_id}", status_code=204)
def discard(import_id: int, db: Session = Depends(get_db), user: User = Depends(CIRCUITS_MANAGE)):
    imp = db.get(CircuitImport, import_id) or _not_found()
    if imp.status != PENDING:
        raise HTTPException(409, "Only a pending import can be discarded; "
                                 "earlier versions are kept as history")
    db.execute(delete(Circuit).where(Circuit.import_id == imp.id))
    db.delete(imp)
    audit(db, user.username, "circuits.import.discard", imp.filename)
    return Response(status_code=204)


def _not_found():
    raise HTTPException(404, "Import not found")
