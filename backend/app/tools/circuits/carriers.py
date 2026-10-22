"""Carrier circuits: the lines from providers (internet, MPLS, point-to-point, SIP...).

Each one has the provider's reference, the contract and the router or firewall interface it
lands on. Its live status is monitoring's last reading of that interface (no device is
contacted from here), so a line that drops is seen within the interface read interval, and
the poller treats these interfaces as watched (a port-down event and email even when the
port has no description). Contracts ending soon are emailed (``core/alerts._carriers``).
"""

import csv
import io
from datetime import date, datetime

from fastapi import APIRouter, Depends, HTTPException, Query, Request, Response
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core import exports
from app.core.audit import audit
from app.core.auth import User, require
from app.core.db import get_db, utcnow
from app.core.models import Device
from app.tools.circuits.importer import _cell, _norm_header, normalize_port
from app.tools.circuits.models import CarrierCircuit

router = APIRouter(prefix="/api/circuits/carriers", tags=["circuits"])

VIEW = require("circuits", "view")
MANAGE = require("circuits", "manage")

SERVICES = ["Internet", "MPLS", "Point-to-point", "SD-WAN underlay", "SIP / voice", "Dark fibre",
            "4G / 5G backup", "Other"]
ENDING_DAYS = 60  # contracts ending within this many days are flagged (and emailed)

# column key -> label, in export / template / import order
COLUMNS = {
    "provider": "Provider", "circuit_ref": "Circuit ID", "service": "Service",
    "bandwidth": "Bandwidth", "site": "Site", "device": "Router or firewall",
    "interface": "Interface", "wan_ip": "WAN IP", "contract_end": "Contract end",
    "cost": "Cost", "support": "Fault line / account", "notes": "Notes",
}
_IMPORT = {_norm_header(v): k for k, v in COLUMNS.items()} | {
    "carrier": "provider", "supplier": "provider", "isp": "provider", "circuit": "circuit_ref",
    "circuit ref": "circuit_ref", "circuit reference": "circuit_ref", "reference": "circuit_ref",
    "speed": "bandwidth", "device": "device", "router": "device", "firewall": "device",
    "port": "interface", "ip": "wan_ip", "ip address": "wan_ip", "end date": "contract_end",
    "renewal": "contract_end", "renewal date": "contract_end", "contract": "contract_end",
    "monthly cost": "cost", "price": "cost", "support": "support", "fault line": "support",
    "type": "service", "location": "site",
}

STATUS_LABELS = {
    "up": "Up", "down": "Down", "admin_down": "Shut down", "device_down": "Router down",
    "not_found": "Interface not found", "not_read": "Not read yet", "not_monitored": "Not monitored",
    "no_link": "No router or interface",
}


class CarrierIn(BaseModel):
    provider: str = Field("", max_length=200)
    circuit_ref: str = Field("", max_length=200)
    service: str = Field("", max_length=100)
    bandwidth: str = Field("", max_length=100)
    site: str = Field("", max_length=100)
    device_id: int | None = None
    interface: str = Field("", max_length=100)
    wan_ip: str = Field("", max_length=100)
    contract_end: date | None = None
    cost: str = Field("", max_length=100)
    support: str = Field("", max_length=255)
    notes: str = Field("", max_length=4000)


# --- live status from monitoring ---------------------------------------------------------

class Live:
    """Monitoring's last reading of every interface and device, loaded once per request."""

    def __init__(self, request: Request | None):
        self.ifaces: dict[int, list[dict]] = {}
        self.states: dict[int, dict] = {}
        self.ok = False
        if request is None:
            return
        try:
            from app.tools.monitoring.store import MonitorStore

            store = MonitorStore(request.app.state.settings.monitor_dir)
            for r in store.interfaces():
                self.ifaces.setdefault(r["device_id"], []).append(r)
            self.states = store.states()
            self.ok = bool(self.states or self.ifaces)
        except Exception:  # noqa: BLE001 - no monitoring data: status "not monitored"
            self.ok = False

    def find(self, device_id: int, interface: str) -> dict | None:
        want = normalize_port(interface)
        for r in self.ifaces.get(device_id, []):
            if want in (normalize_port(r.get("name") or ""), normalize_port(r.get("descr") or "")):
                return r
        return None

    def status(self, c: CarrierCircuit) -> dict:
        out = {"status": "", "in_bps": None, "out_bps": None, "util": None, "speed": None,
               "since": None, "read_at": None}
        if not c.device_id or not c.interface.strip():
            out["status"] = "no_link"
        elif c.device_id not in self.states and c.device_id not in self.ifaces:
            out["status"] = "not_monitored"
        elif (self.states.get(c.device_id) or {}).get("status") == "down":
            out["status"] = "device_down"
        else:
            r = self.find(c.device_id, c.interface)
            if r is None:
                out["status"] = "not_found" if self.ifaces.get(c.device_id) else "not_read"
            else:
                util = [u for u in (r.get("in_util"), r.get("out_util")) if u is not None]
                out.update(status="up" if r.get("oper") == 1 else (
                    "admin_down" if r.get("admin") == 2 else "down"),
                    in_bps=r.get("in_bps"), out_bps=r.get("out_bps"),
                    util=max(util) if util else None, speed=r.get("speed"),
                    since=r.get("since"), read_at=r.get("ts"))
        out["status_label"] = STATUS_LABELS[out["status"]]
        return out


def watched_ports(db: Session) -> dict[int, dict[str, str]]:
    """device id -> {normalised interface: circuit label}: the poller watches these ports."""
    out: dict[int, dict[str, str]] = {}
    for c in db.scalars(select(CarrierCircuit).where(CarrierCircuit.device_id.is_not(None))):
        if c.interface.strip():
            out.setdefault(c.device_id, {})[normalize_port(c.interface)] = label(c)
    return out


def label(c: CarrierCircuit) -> str:
    return " ".join(x for x in (c.provider, c.circuit_ref) if x) or f"circuit {c.id}"


def ending(c: CarrierCircuit, today: date | None = None) -> int | None:
    """Days until the contract ends (negative once it has), or None."""
    if c.contract_end is None:
        return None
    return (c.contract_end - (today or utcnow().date())).days


# --- API ---------------------------------------------------------------------------------

def _out(c: CarrierCircuit, names: dict[int, str], live: Live, sites) -> dict:
    days = ending(c)
    site = sites.resolve(c.site, c.device_id, None)
    return {"id": c.id, "provider": c.provider, "circuit_ref": c.circuit_ref, "service": c.service,
            "bandwidth": c.bandwidth, "site": c.site, "device_id": c.device_id,
            "device": names.get(c.device_id, "") if c.device_id else "", "interface": c.interface,
            "wan_ip": c.wan_ip, "contract_end": c.contract_end.isoformat() if c.contract_end else None,
            "contract_days": days, "contract_ending": days is not None and days <= ENDING_DAYS,
            "cost": c.cost, "support": c.support, "notes": c.notes, "updated_at": c.updated_at,
            "updated_by": c.updated_by, **{k: site[k] for k in ("site_id", "site_name", "site_from")},
            **live.status(c)}


def _all(db: Session, request: Request | None) -> list[dict]:
    from app.tools.circuits.api import SiteIndex

    names = dict(db.execute(select(Device.id, Device.name)).all())
    live, sites = Live(request), SiteIndex(db)
    rows = db.scalars(select(CarrierCircuit).order_by(CarrierCircuit.site, CarrierCircuit.provider,
                                                      CarrierCircuit.circuit_ref))
    return [_out(c, names, live, sites) for c in rows]


@router.get("")
def list_carriers(request: Request, db: Session = Depends(get_db), _: User = Depends(VIEW)):
    from app.core import licence

    return {"circuits": _all(db, request), "services": SERVICES, "ending_days": ENDING_DAYS,
            "monitoring": "monitoring" in licence.status(db)["modules"]}


def _check(db: Session, body: CarrierIn, exclude: int | None = None) -> None:
    if not body.provider.strip() and not body.circuit_ref.strip():
        raise HTTPException(422, "Give the provider or the circuit ID")
    if body.device_id is not None and db.get(Device, body.device_id) is None:
        raise HTTPException(422, "That router or firewall isn't in the inventory")
    if body.circuit_ref.strip():
        q = select(CarrierCircuit).where(CarrierCircuit.circuit_ref == body.circuit_ref.strip(),
                                         CarrierCircuit.provider == body.provider.strip())
        if exclude is not None:
            q = q.where(CarrierCircuit.id != exclude)
        if db.scalar(q) is not None:
            raise HTTPException(409, "That provider already has a circuit with this ID")


def _apply(c: CarrierCircuit, body: CarrierIn, user: str) -> None:
    for k, v in body.model_dump().items():
        setattr(c, k, v.strip() if isinstance(v, str) else v)
    c.updated_at, c.updated_by = utcnow(), user


@router.post("", status_code=201)
def add_carrier(body: CarrierIn, request: Request, db: Session = Depends(get_db),
                user: User = Depends(MANAGE)):
    _check(db, body)
    c = CarrierCircuit()
    _apply(c, body, user.username)
    db.add(c)
    db.flush()
    audit(db, user.username, "circuits.carrier.add", label(c))
    return next(r for r in _all(db, request) if r["id"] == c.id)


def _get(db: Session, circuit_id: int) -> CarrierCircuit:
    c = db.get(CarrierCircuit, circuit_id)
    if c is None:
        raise HTTPException(404, "Carrier circuit not found")
    return c


@router.put("/{circuit_id:int}")
def change_carrier(circuit_id: int, body: CarrierIn, request: Request,
                   db: Session = Depends(get_db), user: User = Depends(MANAGE)):
    c = _get(db, circuit_id)
    _check(db, body, exclude=c.id)
    _apply(c, body, user.username)
    audit(db, user.username, "circuits.carrier.change", label(c))
    return next(r for r in _all(db, request) if r["id"] == c.id)


@router.delete("/{circuit_id:int}", status_code=204)
def delete_carrier(circuit_id: int, db: Session = Depends(get_db), user: User = Depends(MANAGE)):
    c = _get(db, circuit_id)
    db.delete(c)
    audit(db, user.username, "circuits.carrier.delete", label(c))
    return Response(status_code=204)


# --- spreadsheets ------------------------------------------------------------------------

def _cells(r: dict) -> list:
    return [r["provider"], r["circuit_ref"], r["service"], r["bandwidth"], r["site"], r["device"],
            r["interface"], r["wan_ip"], r["contract_end"] or "", r["cost"], r["support"],
            r["notes"], r["status_label"]]


@router.get("/export.csv")
def export_csv(request: Request, db: Session = Depends(get_db), _: User = Depends(VIEW)):
    out = io.StringIO()
    w = exports.csv_writer(out)
    w.writerow([*COLUMNS.values(), "Status"])
    for r in _all(db, request):
        w.writerow(_cells(r))
    return Response("﻿" + out.getvalue(), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": 'attachment; filename="carrier-circuits.csv"'})


@router.get("/template.{fmt}")
def template(fmt: str, _: User = Depends(VIEW)):
    if fmt not in ("csv", "xlsx"):
        raise HTTPException(404, "Not found")
    return exports.template(fmt, list(COLUMNS.values()), "carrier-circuits", notes=[
        "One row per line from a provider. Give the Provider or the Circuit ID (or both).",
        "Router or firewall: the device's name as in the inventory; Interface: as the "
        "device names it (e.g. GigabitEthernet0/0, ge-0/0/1). Their status then comes from "
        "monitoring.",
        "Contract end: a date (YYYY-MM-DD). Ends within 60 days are flagged and emailed.",
        "Uploading adds new circuits and updates those with the same Provider and Circuit ID; "
        "nothing is removed.",
    ], choices={"Service": SERVICES}, sheet="Carrier circuits")


def _rows_of(data: bytes, name: str) -> list[dict[str, str]]:
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        from app.tools.circuits.importer import _check_unpacked_size

        _check_unpacked_size(data)
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:  # noqa: BLE001
            raise HTTPException(422, "Not a readable .xlsx workbook") from None
        try:
            table = [tuple(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        finally:
            wb.close()
    elif name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1252", errors="replace")
        first = next((ln for ln in text.splitlines() if ln.strip()), "")
        table = [tuple(r) for r in csv.reader(io.StringIO(text),
                                              delimiter=max(",;\t", key=first.count) if first else ",")]
    else:
        raise HTTPException(422, "Upload an Excel workbook (.xlsx) or a CSV file")
    for idx, row in enumerate(table[:15]):
        cols = {i: _IMPORT[_norm_header(c)] for i, c in enumerate(row) if _norm_header(c) in _IMPORT}
        if "provider" in cols.values() or "circuit_ref" in cols.values():
            if len(cols) >= 2:
                out = []
                for raw in table[idx + 1:idx + 1 + 5000]:
                    rec = {k: _value(raw[i]) if i < len(raw) else "" for i, k in cols.items()}
                    if any(rec.values()):
                        out.append(rec)
                return out
    raise HTTPException(422, "No header row with 'Provider' or 'Circuit ID' in the first 15 rows "
                             "(download the blank template to see the column names)")


def _value(v) -> str:
    """A cell as text; a date cell (Excel) as YYYY-MM-DD."""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return _cell(v)


def _date(text: str) -> date | None:
    text = (text or "").strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"'{text}' is not a date (use YYYY-MM-DD)")


@router.put("/upload")
async def upload(request: Request, filename: str = Query(..., max_length=255),
                 user: User = Depends(MANAGE)):
    """Add and update carrier circuits from a spreadsheet: a row with the same Provider and
    Circuit ID as an existing one updates it (empty cells keep what is there); others are
    added. Nothing is removed."""
    from app.core.db import session_scope

    data = await request.body()
    if len(data) > request.app.state.settings.circuits_max_upload_mb * 2**20:
        raise HTTPException(413, "File too large")
    rows = _rows_of(data, filename.lower())
    added = updated = 0
    problems: list[str] = []
    with session_scope() as db:
        devices = {d.name.lower(): d.id for d in db.scalars(select(Device))}
        existing = {(c.provider.lower(), c.circuit_ref.lower()): c
                    for c in db.scalars(select(CarrierCircuit)) if c.circuit_ref}
        for n, rec in enumerate(rows, 2):
            try:
                end = _date(rec.get("contract_end", ""))
            except ValueError as exc:
                problems.append(f"Row {n}: {exc}")
                continue
            dev = rec.get("device", "").strip()
            dev_id = devices.get(dev.lower()) if dev else None
            if dev and dev_id is None:
                problems.append(f"Row {n}: '{dev}' isn't in the inventory (left unlinked)")
            if not rec.get("provider") and not rec.get("circuit_ref"):
                problems.append(f"Row {n}: no Provider or Circuit ID (skipped)")
                continue
            key = (rec.get("provider", "").strip().lower(), rec.get("circuit_ref", "").strip().lower())
            c = existing.get(key) if key[1] else None
            if c is None:
                c = CarrierCircuit()
                db.add(c)
                if key[1]:
                    existing[key] = c
                added += 1
            else:
                updated += 1
            for k, v in rec.items():
                if k in ("device", "contract_end") or not v.strip():
                    continue
                setattr(c, k, v.strip()[:4000 if k == "notes" else 255])
            if dev_id is not None:
                c.device_id = dev_id
            if end is not None:
                c.contract_end = end
            c.updated_at, c.updated_by = utcnow(), user.username
        audit(db, user.username, "circuits.carrier.upload",
              f"{filename}: {added} added, {updated} updated, {len(problems)} problems")
    return {"added": added, "updated": updated, "problems": problems[:100]}


def site_counts(db: Session) -> dict[int, int]:
    """Carrier circuits per site (site id -> count), for the Sites page."""
    from app.tools.circuits.api import SiteIndex

    sites, out = SiteIndex(db), {}
    for c in db.scalars(select(CarrierCircuit)):
        s = sites.resolve(c.site, c.device_id, None)["site_id"]
        if s:
            out[s] = out.get(s, 0) + 1
    return out


def contract_alert_rows(db: Session, today: date | None = None) -> list[tuple[CarrierCircuit, int]]:
    """Circuits whose contract ends within ENDING_DAYS (or has ended in the last 30 days)."""
    out = []
    for c in db.scalars(select(CarrierCircuit).where(CarrierCircuit.contract_end.is_not(None))):
        d = ending(c, today)
        if d is not None and -30 <= d <= ENDING_DAYS:
            out.append((c, d))
    return sorted(out, key=lambda x: x[1])

