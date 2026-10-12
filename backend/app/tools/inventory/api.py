"""HTTP API for the Inventory page: the rows, the columns (built-in and the ones admins add),
'Manage this device', and import and export."""

import csv
import datetime as dt
import io
import json

from fastapi import APIRouter, Depends, HTTPException, Query, Request, Response
from fastapi.concurrency import run_in_threadpool
from pydantic import BaseModel
from pydantic import Field as PField
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core import exports, licence
from app.core.audit import audit
from app.core.auth import User, current_user, require, require_admin
from app.core.db import get_db
from app.core.models import Credential, Device
from app.core.platforms import PLATFORMS
from app.tools.inventory import fields as fl
from app.tools.inventory.models import Asset, AssetField

router = APIRouter(prefix="/api/inventory", tags=["inventory"])
MANAGE = require("inventory", "manage")
ADDRESSES = require("inventory", "addresses")
MAX_IMPORT = 10 * 2**20
MAX_ROWS = 20000


# --- rows -------------------------------------------------------------------------------

def _plain(value):
    return value.isoformat() if isinstance(value, dt.date) else value


def _status_of_devices(db: Session, ids: list[int]) -> tuple[dict, dict, dict]:
    """Backup state, firmware facts and compliance for managed rows (if those tools are on)."""
    modules = set(licence.status(db)["modules"])
    backups, facts, comp = {}, {}, {}
    if not ids:
        return backups, facts, comp
    from app.tools.firmware_upgrade.facts import best_standard, compliance
    from app.tools.firmware_upgrade.models import DeviceFacts, FirmwareStandard

    facts = {f.device_id: f for f in db.scalars(select(DeviceFacts).where(DeviceFacts.device_id.in_(ids)))}
    if "backups" in modules:
        from app.tools.config_backup.models import BackupState

        backups = {s.device_id: s.last_status for s in
                   db.scalars(select(BackupState).where(BackupState.device_id.in_(ids)))}
    if "firmware" in modules:
        standards = list(db.scalars(select(FirmwareStandard)))
        for d in db.scalars(select(Device).where(Device.id.in_(ids))):
            f = facts.get(d.id)
            std = best_standard(standards, d.platform, f.model if f else None)
            comp[d.id] = compliance(f.version if f else None, std.target_version if std else None)
    return backups, facts, comp


def rows(db: Session) -> list[dict]:
    fl.sync_devices(db)
    custom = fl.custom_fields(db)
    assets = list(db.scalars(select(Asset).order_by(func.lower(Asset.site), func.lower(Asset.name))))
    ids = [a.device_id for a in assets if a.device_id]
    devices = {d.id: d for d in db.scalars(select(Device).where(Device.id.in_(ids)))} if ids else {}
    backups, facts, comp = _status_of_devices(db, ids)
    out = []
    for a in assets:
        values = {f.key: _plain(getattr(a, f.key)) for f in fl.STANDARD}
        stored = json.loads(a.custom or "{}")
        values |= {f.key: stored.get(f.key[2:]) for f in custom}
        auto = {}
        device = devices.get(a.device_id)
        if device:
            known = fl.device_values(device, facts.get(device.id))
            auto = {f.key: known.get(f.auto) for f in fl.STANDARD if f.auto and known.get(f.auto)}
        out.append({"id": a.id, "managed": bool(device), "values": values, "auto": auto,
                    "device": {"id": device.id, "name": device.name, "platform": device.platform}
                    if device else None,
                    "backup": backups.get(a.device_id), "compliance": comp.get(a.device_id),
                    "updated_at": a.updated_at, "updated_by": a.updated_by})
    return out


def _effective(row: dict, key: str):
    value = row["values"].get(key)
    return value if value not in (None, "") else row["auto"].get(key, "")


@router.get("/meta")
def meta(db: Session = Depends(get_db), user: User = Depends(current_user)):
    return {"fields": fl.describe(fl.all_fields(db)), "kinds": fl.KINDS,
            "can_manage": user.can("inventory", "manage"),
            "can_add_devices": user.can("inventory", "addresses"), "is_admin": user.is_admin}


@router.get("/assets")
def list_assets(db: Session = Depends(get_db), _: User = Depends(current_user)):
    return rows(db)


class AssetIn(BaseModel):
    values: dict = PField(default_factory=dict)


def _get(db: Session, asset_id: int) -> Asset:
    asset = db.get(Asset, asset_id)
    if asset is None:
        raise HTTPException(404, "Not found in the inventory")
    return asset


def _set(db: Session, asset: Asset, values: dict, user: User) -> list[str]:
    try:
        changed = fl.apply(asset, values, fl.all_fields(db))
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from None
    asset.updated_by = user.username
    return changed


@router.post("/assets", status_code=201)
def add_asset(body: AssetIn, db: Session = Depends(get_db), user: User = Depends(MANAGE)):
    asset = Asset(name="", custom="{}")
    _set(db, asset, {"status": "in service", **body.values}, user)
    if not asset.name.strip():
        raise HTTPException(422, "Name: required")
    db.add(asset)
    db.flush()
    audit(db, user.username, "inventory.add", f"{asset.name} ({asset.type or 'no type'})")
    return {"id": asset.id}


@router.put("/assets/{asset_id}")
def update_asset(asset_id: int, body: AssetIn, db: Session = Depends(get_db),
                 user: User = Depends(MANAGE)):
    asset = _get(db, asset_id)
    if asset.device_id and "name" in body.values and \
            str(body.values["name"]).strip() != asset.name:
        raise HTTPException(422, "A managed device is renamed on Config backups (Edit), so the "
                                 "tools and the inventory keep the same name")
    changed = _set(db, asset, body.values, user)
    if changed:
        audit(db, user.username, "inventory.update", f"{asset.name}: {', '.join(changed)}")
    return {"id": asset.id, "changed": changed}


@router.delete("/assets/{asset_id}", status_code=204)
def delete_asset(asset_id: int, db: Session = Depends(get_db), user: User = Depends(MANAGE)):
    asset = _get(db, asset_id)
    if asset.device_id:
        raise HTTPException(409, "This device is managed by the tools: delete it on Config "
                                 "backups first (its inventory details then stay here)")
    db.delete(asset)
    audit(db, user.username, "inventory.delete", asset.name)
    return Response(status_code=204)


class ManageIn(BaseModel):
    platform: str
    address: str = PField(min_length=1, max_length=255)
    credential_id: int | None = None


@router.post("/assets/{asset_id}/manage")
def manage(asset_id: int, body: ManageIn, db: Session = Depends(get_db),
           user: User = Depends(ADDRESSES)):
    """Turn an inventory-only row into a device the tools manage (counts for the licence)."""
    asset = _get(db, asset_id)
    if asset.device_id:
        raise HTTPException(409, "Already managed")
    if body.platform not in PLATFORMS:
        raise HTTPException(422, "Choose the device's platform")
    if body.credential_id is not None and db.get(Credential, body.credential_id) is None:
        raise HTTPException(422, "Unknown credential profile")
    name = asset.name.strip()[:100]
    if db.scalar(select(Device.id).where(func.lower(Device.name) == name.lower())):
        raise HTTPException(409, f"A device called {name} already exists")
    licence.room_for(db)
    device = Device(name=name, address=body.address.strip(),
                    platform=body.platform, site=asset.site, credential_id=body.credential_id)
    db.add(device)
    try:
        db.flush()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, f"A device called {asset.name} already exists") from None
    asset.device_id = device.id
    if not asset.mgmt_ip:
        asset.mgmt_ip = device.address
    audit(db, user.username, "device.create", f"{device.name} ({device.address}) from the inventory")
    return {"device_id": device.id}


# --- the columns admins add ------------------------------------------------------------------

class FieldIn(BaseModel):
    label: str = PField(min_length=1, max_length=100)
    kind: str = "text"
    choices: list[str] = PField(default_factory=list, max_length=100)


def _check_field(db: Session, body: FieldIn, exclude: int | None = None) -> None:
    if body.kind not in fl.KINDS:
        raise HTTPException(422, f"Kind: one of {', '.join(fl.KINDS)}")
    if body.kind == "choice" and not [c for c in body.choices if c.strip()]:
        raise HTTPException(422, "A pick-list needs at least one choice")
    used = {f.label.lower() for f in fl.STANDARD} | {
        c.label.lower() for c in db.scalars(select(AssetField)) if c.id != exclude}
    if body.label.strip().lower() in used:
        raise HTTPException(422, f"There is already a column called {body.label.strip()}")


@router.post("/fields", status_code=201)
def add_field(body: FieldIn, db: Session = Depends(get_db), user: User = Depends(require_admin)):
    _check_field(db, body)
    key, n = fl.slug(body.label), 2
    base = key
    while db.scalar(select(AssetField.id).where(AssetField.key == key)):
        key, n = f"{base}_{n}", n + 1
    position = (db.scalar(select(func.max(AssetField.position))) or 0) + 1
    f = AssetField(key=key, label=body.label.strip(), kind=body.kind, position=position,
                   choices=json.dumps([c.strip() for c in body.choices if c.strip()]))
    db.add(f)
    db.flush()
    audit(db, user.username, "inventory.column_add", f"{f.label} ({f.kind})")
    return {"key": f"x_{f.key}"}


def _field(db: Session, key: str) -> AssetField:
    f = db.scalar(select(AssetField).where(AssetField.key == key.removeprefix("x_")))
    if f is None:
        raise HTTPException(404, "No such column")
    return f


@router.put("/fields/{key}")
def change_field(key: str, body: FieldIn, db: Session = Depends(get_db),
                 user: User = Depends(require_admin)):
    f = _field(db, key)
    if body.label.strip().lower() != f.label.lower():
        _check_field(db, body, exclude=f.id)
    if body.kind != f.kind:
        raise HTTPException(422, "A column's kind can't change: add a new column instead")
    old = f.label
    f.label = body.label.strip()
    f.choices = json.dumps([c.strip() for c in body.choices if c.strip()])
    audit(db, user.username, "inventory.column_change", f"{old} -> {f.label}")
    return {"key": f"x_{f.key}"}


@router.delete("/fields/{key}", status_code=204)
def delete_field(key: str, db: Session = Depends(get_db), user: User = Depends(require_admin)):
    f = _field(db, key)
    for asset in db.scalars(select(Asset).where(Asset.custom.like(f'%"{f.key}"%'))):
        values = json.loads(asset.custom or "{}")
        values.pop(f.key, None)
        asset.custom = json.dumps(values)
    db.delete(f)
    audit(db, user.username, "inventory.column_delete", f.label)
    return Response(status_code=204)


# --- export and import -------------------------------------------------------------------------

def _table(db: Session) -> tuple[list[str], list[list]]:
    fields = fl.all_fields(db)
    head = [f.label for f in fields] + ["Managed"]
    body = [[_effective(r, f.key) for f in fields] + ["yes" if r["managed"] else "no"]
            for r in rows(db)]
    return head, body


@router.get("/export.csv")
def export_csv(db: Session = Depends(get_db), _: User = Depends(current_user)):
    head, body = _table(db)
    out = io.StringIO()
    w = exports.csv_writer(out)
    w.writerow(head)
    for row in body:
        w.writerow(["" if v is None else v for v in row])
    return Response(out.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": 'attachment; filename="inventory.csv"'})


@router.get("/export.xlsx")
def export_xlsx(db: Session = Depends(get_db), _: User = Depends(current_user)):
    from openpyxl import Workbook

    head, body = _table(db)
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"
    ws.append(head)
    for row in body:
        exports.append_row(ws, ["" if v is None else v for v in row])
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), headers={"Content-Disposition": 'attachment; filename="inventory.xlsx"'},
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _read_table(data: bytes, filename: str) -> list[list[str]]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        from app.tools.circuits.importer import CircuitImportError, _check_unpacked_size

        try:
            _check_unpacked_size(data)
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except CircuitImportError as exc:
            raise HTTPException(422, str(exc)) from None
        except Exception as exc:  # noqa: BLE001 - not a workbook
            raise HTTPException(422, f"Not a readable workbook ({type(exc).__name__})") from None
        try:
            return [["" if c is None else c for c in r] for r in
                    wb.worksheets[0].iter_rows(values_only=True, max_row=MAX_ROWS + 1)]
        finally:
            wb.close()
    if filename.lower().endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        return [r for i, r in enumerate(csv.reader(io.StringIO(text))) if i <= MAX_ROWS]
    raise HTTPException(422, "Upload an Excel workbook (.xlsx) or a CSV file")


def plan_import(db: Session, table: list[list]) -> dict:
    """What an import would do. Rows are matched by serial number, then asset tag, then name."""
    fl.sync_devices(db)  # devices added since the page was last opened
    fields = fl.all_fields(db)
    names = {}
    for f in fields:
        names[f.label.lower()] = f
        names[f.key.lower()] = f
    if not table or not any(str(c).strip() for c in table[0]):
        raise HTTPException(422, "The first row must hold the column names")
    header = [str(c).strip() for c in table[0]]
    columns = {i: names[h.lower()] for i, h in enumerate(header) if h.lower() in names}
    unknown = [h for i, h in enumerate(header) if h and i not in columns and h.lower() != "managed"]
    if not any(f.key == "name" for f in columns.values()):
        raise HTTPException(422, "A 'Name' column is needed")
    existing = list(db.scalars(select(Asset)))
    by = {k: {getattr(a, k).strip().lower(): a for a in existing if getattr(a, k).strip()}
          for k in ("serial", "asset_tag", "name")}
    new, updated, errors, seen = [], [], [], set()
    for line, raw in enumerate(table[1:], start=2):
        values = {f.key: raw[i] if i < len(raw) else "" for i, f in columns.items()}
        if not any(str(v).strip() for v in values.values()):
            continue
        values = {k: v for k, v in values.items() if str(v).strip() != ""}
        match = next((by[k].get(str(values[k]).strip().lower()) for k in ("serial", "asset_tag", "name")
                      if values.get(k) and by[k].get(str(values[k]).strip().lower())), None)
        if match and match.device_id:
            values.pop("name", None)  # a managed device keeps the name the tools use
        probe = Asset(name=match.name if match else "", custom=match.custom if match else "{}")
        if match:
            for f in fl.STANDARD:
                setattr(probe, f.key, getattr(match, f.key))
        try:
            changed = fl.apply(probe, values, fields)
        except ValueError as exc:
            errors.append({"row": line, "message": str(exc)})
            continue
        if not match and not probe.name.strip():
            errors.append({"row": line, "message": "Name: required for a new item"})
            continue
        key = id(match) if match else probe.name.lower()
        if key in seen:
            errors.append({"row": line, "message": f"{probe.name} is in the file twice"})
            continue
        seen.add(key)
        if match:
            if changed:
                updated.append({"row": line, "id": match.id, "name": match.name, "changes": changed,
                                "values": values})
        else:
            new.append({"row": line, "name": probe.name, "values": values})
    return {"new": new, "updated": updated, "errors": errors, "unknown_columns": unknown,
            "columns": [f.label for f in columns.values()]}


@router.put("/import")
async def import_file(request: Request, filename: str = Query(..., max_length=255),
                      commit: bool = False, user: User = Depends(MANAGE)):
    """Upload a spreadsheet (raw body). Without ``commit`` it only shows what would change;
    with it, the changes are made (refused while the file has errors)."""
    body = bytearray()
    async for chunk in request.stream():
        body += chunk
        if len(body) > MAX_IMPORT:
            raise HTTPException(413, f"File is larger than {MAX_IMPORT // 2**20} MB")
    table = await run_in_threadpool(_read_table, bytes(body), filename)

    def work():
        from app.core.db import session_scope

        with session_scope() as db:
            plan = plan_import(db, table)
            if not commit:
                return plan
            if plan["errors"]:
                raise HTTPException(422, f"Fix the {len(plan['errors'])} row(s) with errors first")
            fields = fl.all_fields(db)
            for item in plan["new"]:
                asset = Asset(name="", custom="{}", updated_by=user.username)
                fl.apply(asset, {"status": "in service", **item["values"]}, fields)
                db.add(asset)
            for item in plan["updated"]:
                asset = db.get(Asset, item["id"])
                fl.apply(asset, item["values"], fields)
                asset.updated_by = user.username
            db.flush()
            fl.sync_devices(db)
            audit(db, user.username, "inventory.import",
                  f"{filename}: {len(plan['new'])} new, {len(plan['updated'])} updated")
            return {**plan, "committed": True}

    return await run_in_threadpool(work)
