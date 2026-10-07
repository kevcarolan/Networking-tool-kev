"""Live circuit locations: MAC/ARP parsing, uplink rules, resolving, the scan and the API."""

import io
from datetime import datetime

import pytest
from openpyxl import load_workbook

from app.tools.circuits import live
from app.tools.circuits.live import Seen, resolve
from app.tools.circuits.scan import last_slot, parse_schedule
from tests import firmware_samples as s
from tests.test_circuits import ROWS, _upload, make_xlsx


# --- parsers ------------------------------------------------------------------

def test_ios_mac_table_skips_cpu_and_duplicates():
    entries = live.parse_mac_table(s.IOS_MAC_TABLE)
    got = {(e.mac, e.vlan, e.port_norm) for e in entries}
    assert ("02:00:5e:10:c0:2f", "34", "gi2/0/13") in got
    assert ("02:00:5e:30:62:92", "36", "gi2/0/15") in got  # sticky (STATIC) MACs count
    assert not any(e.mac.startswith("01:") or e.mac == "ff:ff:ff:ff:ff:ff" for e in entries)
    assert len(entries) == 7


def test_nxos_mac_table_skips_peer_link_and_sup():
    entries = live.parse_mac_table(s.NXOS_MAC_TABLE)
    assert [(e.mac, e.vlan, e.port) for e in entries] == [
        ("00:50:56:a1:00:01", "10", "Eth1/5"), ("00:50:56:a1:00:02", "10", "Po10")]


def test_awplus_mac_table_port_before_mac():
    entries = live.parse_mac_table(s.AWPLUS_MAC_TABLE)
    assert [(e.vlan, e.port) for e in entries] == [("20", "port1.0.1"), ("20", "port1.0.24")]


@pytest.mark.parametrize("out, pair", [
    (s.IOS_ARP, ("10.20.66.16", "02:00:5e:20:22:69")),
    (s.NXOS_ARP, ("10.10.0.6", "00:50:56:a1:00:02")),
    (s.AWPLUS_ARP, ("10.20.0.5", "00:00:cd:28:07:52")),
    (s.ASA_ARP, ("10.20.34.70", "02:00:5e:10:c0:aa")),
])
def test_arp(out, pair):
    arps = live.parse_arp(out)
    assert pair in arps
    assert all(ip != "10.20.66.99" for ip, _ in arps)  # incomplete entries left out


def test_uplinks_from_cdp_lldp_and_port_channels():
    assert live.parse_uplinks_cdp(s.CDP_DETAIL) == {"gi1/0/1", "gi1/1/1"}
    assert live.parse_uplinks_lldp(s.LLDP_DETAIL_PHONE) == {"gi1/0/1"}  # a phone isn't an uplink
    assert live.parse_uplinks_lldp(s.AWPLUS_LLDP) == {"port1.0.24"}
    assert live.parse_port_channels(s.ETHERCHANNEL) == {"po1", "gi1/1/1", "gi2/1/1"}
    assert live.parse_port_channels(s.NXOS_PORT_CHANNEL) == {"po10", "eth1/49", "eth1/50"}


def test_read_device_marks_uplinks():
    r = live.read_device({"mac": s.IOS_MAC_TABLE, "arp": s.IOS_ARP, "cdp": s.CDP_DETAIL,
                          "lldp": s.LLDP_DETAIL_PHONE, "pc": s.ETHERCHANNEL})
    edge = {e.port_norm for e in r.edge}
    assert edge == {"gi2/0/13", "gi2/0/14", "gi2/0/15", "gi1/0/3"}  # not Gi1/0/1, Gi1/1/1, Po1
    assert len(r.arps) == 3
    asa = live.read_device({"mac": s.INVALID, "arp": s.ASA_ARP})
    assert asa.macs is None and len(asa.arps) == 2


# --- resolving ----------------------------------------------------------------

EDGES = {"aa:aa:aa:aa:aa:01": [Seen(1, "Gi2/0/13", "gi2/0/13", "34")],
         "aa:aa:aa:aa:aa:02": [Seen(1, "Gi2/0/13", "gi2/0/13", "34"),
                               Seen(2, "Gi1/0/2", "gi1/0/2", "34")]}


@pytest.mark.parametrize("mac, ip, documented, status", [
    ("aa:aa:aa:aa:aa:01", "", (1, "gi2/0/13"), live.MATCHES),
    ("aa:aa:aa:aa:aa:01", "", (1, "gi2/0/12"), live.MOVED),
    ("aa:aa:aa:aa:aa:01", "", (2, "gi2/0/13"), live.MOVED),
    ("aa:aa:aa:aa:aa:01", "", (None, ""), live.FOUND),
    ("aa:aa:aa:aa:aa:02", "", (1, "gi2/0/13"), live.SEVERAL),
    ("aa:aa:aa:aa:aa:09", "", (1, "gi2/0/13"), live.NOT_SEEN),
    ("", "10.0.0.1", (1, "gi2/0/13"), live.MATCHES),  # MAC from ARP
    ("", "10.0.0.2", (1, "gi2/0/13"), live.NOT_SEEN),
    ("", "", (1, "gi2/0/13"), live.NO_ADDRESS),
])
def test_resolve(mac, ip, documented, status):
    r = resolve(mac, ip, documented, EDGES, {"10.0.0.1": "aa:aa:aa:aa:aa:01"})
    assert r["live_status"] == status
    if not mac and ip == "10.0.0.1":
        assert r["live_mac_source"] == "arp"


# --- schedule -----------------------------------------------------------------

def test_schedule():
    assert parse_schedule("") is None
    assert parse_schedule("sun 02:00") == (6, 2, 0)
    assert parse_schedule("Sunday 02:30") == (6, 2, 30)
    assert parse_schedule("daily 23:15") == (None, 23, 15)
    with pytest.raises(ValueError):
        parse_schedule("every sunday")
    wed = datetime(2026, 10, 7, 12, 0)  # a Wednesday
    assert last_slot((6, 2, 0), wed) == datetime(2026, 10, 4, 2, 0)
    assert last_slot((None, 2, 0), wed) == datetime(2026, 10, 7, 2, 0)
    assert last_slot((2, 13, 0), wed) == datetime(2026, 9, 30, 13, 0)


def test_scheduled_scan_runs_once_per_slot(app, admin, credential, fw_collector):
    svc = app.state.circuit_scan
    svc.schedule = parse_schedule("sun 02:00")
    svc.clock = lambda: datetime(2026, 10, 4, 2, 5)
    assert svc.due_trigger() == "schedule:2026-10-04 02:00"
    assert svc.tick() is True
    assert svc.due_trigger() is None  # already done
    svc.clock = lambda: datetime(2026, 10, 4, 9, 0)  # too long after the slot: wait a week
    assert svc.due_trigger() is None


# --- API end to end -----------------------------------------------------------

@pytest.fixture
def switch(admin, credential):
    r = admin.post("/api/devices", json={"name": "SITEA-ACC-SW01", "address": "10.20.0.11",
                                         "platform": "cisco_ios", "site": "SITEA",
                                         "credential_id": credential["id"]})
    return r.json()


def _scan(app, fw_collector):
    fw_collector.by_command["SITEA-ACC-SW01"] = {
        "show mac address-table": s.IOS_MAC_TABLE, "show ip arp": s.IOS_ARP,
        "show cdp neighbors detail": s.CDP_DETAIL, "show lldp neighbors detail": s.LLDP_DETAIL,
        "show etherchannel summary": s.ETHERCHANNEL}
    return app.state.circuit_scan.scan("manual:test")


def test_scan_and_live_status(app, admin, switch, fw_collector):
    rows = [list(r) for r in ROWS]
    rows[2][6] = None  # BMS row: no MAC, found through ARP (10.20.66.16)
    imp = _upload(admin, make_xlsx(rows=rows)).json()
    admin.post(f"/api/circuits/imports/{imp['id']}/commit")
    before = admin.get("/api/circuits").json()
    assert before["circuits"][0]["live_status_label"] == "Not scanned yet"

    scan_id = _scan(app, fw_collector)
    data = admin.get("/api/circuits").json()
    assert data["scan"]["last"]["id"] == scan_id and data["scan"]["last"]["devices_ok"] == 1
    by_name = {c["device_name"]: c for c in data["circuits"]}
    assert by_name["DC-SITEA-GF-03"]["live_status"] == live.MATCHES
    moved = by_name["DC-SITEA-GF-02"]
    assert moved["live_status"] == live.MOVED and moved["live_port"] == "Gi2/0/14"
    assert moved["live_switch"] == "SITEA-ACC-SW01"
    bms = by_name["BMS_CP03"]
    assert bms["live_status"] == live.MATCHES and bms["live_mac_source"] == "arp"
    # learned only on the uplink to ACC-SW02 (not in NetOps): not located
    assert by_name["IC-SITEA-GATE-01"]["live_status"] == live.NOT_SEEN
    assert by_name["IC-SITEA-GF-04"]["live_status"] == live.MATCHES

    fd = admin.get(f"/api/circuits/for-device/{switch['id']}").json()
    assert fd["count"] == 4

    csv = admin.get("/api/circuits/export.csv").text
    assert "Live status" in csv.splitlines()[0] and "Moved" in csv

    wb = load_workbook(io.BytesIO(admin.get("/api/circuits/export-updated.xlsx").content))
    ws = wb.active
    header = [c.value for c in ws[1]]
    row = {header[i]: c.value for i, c in enumerate(next(
        r for r in ws.iter_rows(min_row=2) if r[header.index("Device Name")].value == "DC-SITEA-GF-02"))}
    assert row["PORT"] == "Gi2/0/14" and row["Documented PORT"] == "Gi2/0/12"
    assert row["DEVICE"] == "SITEA-ACC-SW01" and row["Live status"] == "Moved"


def test_last_seen_is_kept_when_a_mac_disappears(app, admin, switch, fw_collector):
    imp = _upload(admin, make_xlsx()).json()
    admin.post(f"/api/circuits/imports/{imp['id']}/commit")
    _scan(app, fw_collector)
    fw_collector.by_command["SITEA-ACC-SW01"]["show mac address-table"] = \
        s.IOS_MAC_TABLE.replace("0200.5e10.c02f", "0200.5e10.ffff")
    app.state.circuit_scan.scan("manual:test")
    c = admin.get("/api/circuits").json()["circuits"][0]
    assert c["live_status"] == live.NOT_SEEN
    assert c["last_seen_at"] and c["last_seen_where"] == "SITEA-ACC-SW01 Gi2/0/13"


def test_scan_failure_is_recorded(app, admin, switch, fw_collector):
    from app.core.ssh import UNREACHABLE, DeviceError

    fw_collector.errors["SITEA-ACC-SW01"] = DeviceError(UNREACHABLE, "unreachable")
    app.state.circuit_scan.scan("manual:test")
    scans = admin.get("/api/circuits/scans").json()
    assert scans[0]["devices_failed"] == 1 and "unreachable" in scans[0]["errors"]


def test_scan_now_is_admin_only(app, client, admin, switch):
    app.state.circuit_scan.scan_in_background = lambda trigger, ids=None: True
    assert admin.post("/api/circuits/scan", json={}).status_code == 202
    assert admin.post("/api/circuits/scan", json={"device_id": 999}).status_code == 404
    from app.core.auth import VIEWER, User

    app.state.authenticator.authenticate = lambda u, p: User(username=u, role=VIEWER)
    client.post("/api/auth/logout")
    client.post("/api/auth/login", json={"username": "v", "password": "x"})
    assert client.post("/api/circuits/scan", json={}).status_code == 403
