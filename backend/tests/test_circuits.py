import io

import pytest
from openpyxl import Workbook

from app.tools.circuits.importer import (CircuitImportError, match_switches, normalize_mac,
                                         normalize_port, parse_workbook, switch_matches)

BANDS = ["SERVICE", "", "DEVICE INFORMATION", "", "", "", "", "LOCATION", "", "IP DETAILS", "",
         "", "NETWORK PORT", "", "", "", "DRAWING DETAILS", ""]
HEADER = ["IP Partition", "Service", "ENGineering prefix", "Device Name", "Device Type",
          "Manufacturer", "Hardware Mac Address", "Floor Level", "Room / Area", "IP Address",
          "Subnet Mask", "Default Gateway", "VLAN", "TYPE", "PORT", "DEVICE", "Drawing Number",
          "Device Drawing Reference"]
ROWS = [
    ["SEC", "ACS", None, "DC-SITEA-GF-03", "DoorCtl 400", "AcmeSec", "02005e10c02f", "Ground Floor",
     "SEC01", "10.20.34.62", "255.255.254.0", "10.20.34.1", 34, "STATIC", "Gi2/0/13", "ACC-SW01",
     "SITEA-SEC-GF-DR-1001", "Intelligent Door Controller"],
    ["SEC", "ACS", None, "DC-SITEA-GF-02", "DoorCtl 400", "AcmeSec", "02005e10c05e", "Ground Floor",
     "SEC01", "10.20.34.61", "255.255.254.0", "10.20.34.1", 34, "STATIC", "Gi2/0/12", "ACC-SW01",
     "SITEA-SEC-GF-DR-1001", "Intelligent Door Controller"],
    ["BMS Co", "BMS", "SITEA_DC1_00_N_RM11_BMS_CP03", "BMS_CP03", None, None,
     "02005E202269", None, None, "10.20.66.16", "255.255.255.0", "10.20.66.1", 66, "STATIC",
     "Gi1/0/3", "ACC-SW01", None, None],
    ["SEC", "Intercom", None, "IC-SITEA-GATE-01", "IP Intercom 1", "VendorX",
     "02005E306248", "N/A", "Vehicle Gate", "10.20.36.34", "255.255.255.0", "10.20.36.1", 36,
     "STATIC", "Gi2/0/8", "ACC-SW02", "SITEA-SEC-EP-DR-1002", "IC-SITEA-GATE-01"],
    ["SEC", "Intercom", None, "IC-SITEA-GF-04", "IP Intercom 1", "VendorX",
     "02005E306292", "Ground Floor", "Turnstile", "10.20.36.39", "255.255.255.0", "10.20.36.1",
     36, "STATIC", "Gi2/0/15", "ACC-SW01", None, "IC-SITEA-GF-04"],
]


def make_xlsx(rows=ROWS, header=HEADER, bands=BANDS, extra_sheet_first=False) -> bytes:
    wb = Workbook()
    ws = wb.active
    if extra_sheet_first:
        ws.title = "Notes"
        ws.append(["This sheet has no header"])
        ws = wb.create_sheet("Circuits")
    else:
        ws.title = "Circuits"
    ws.append([None] * len(header))  # row 1: empty
    ws.append(bands)                 # row 2: band names
    ws.append(header)                # row 3: column names
    for r in rows:
        ws.append(r)
    ws.append([None] * len(header))  # trailing empty row is ignored
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- parsing and normalisation ----------------------------------------------

def test_parse_finds_header_under_bands():
    parsed = parse_workbook(make_xlsx(extra_sheet_first=True))
    assert parsed.sheet == "Circuits" and parsed.header_row == 3
    assert len(parsed.rows) == 5 and parsed.rows[0].row_no == 4
    r = parsed.rows[0]
    assert r.values["device_name"] == "DC-SITEA-GF-03"
    assert r.values["vlan"] == "34" and r.values["switch"] == "ACC-SW01"
    assert r.port_norm == "gi2/0/13" and r.switch_norm == "acc-sw01"
    assert r.mac_norm == "02:00:5e:10:c0:2f" and r.warnings == []
    assert parsed.rows[2].values["eng_prefix"] == "SITEA_DC1_00_N_RM11_BMS_CP03"


def test_parse_column_order_does_not_matter():
    order = list(reversed(range(len(HEADER))))
    parsed = parse_workbook(make_xlsx(rows=[[r[i] for i in order] for r in ROWS],
                                      header=[HEADER[i] for i in order], bands=[]))
    assert parsed.rows[3].values["switch"] == "ACC-SW02"
    assert parsed.rows[3].values["service"] == "Intercom"


def test_parse_warnings():
    bad = list(ROWS[0])
    bad[9], bad[10], bad[12], bad[6], bad[14] = "10.20.34.999", "255.0.255.0", 5000, "xyz", None
    parsed = parse_workbook(make_xlsx(rows=[bad]))
    w = " | ".join(parsed.rows[0].warnings)
    for text in ("only one of", "IP address", "subnet mask", "VLAN", "MAC"):
        assert text in w
    nothing = list(ROWS[0])
    nothing[6], nothing[9], nothing[14], nothing[15] = None, None, None, None
    w = " | ".join(parse_workbook(make_xlsx(rows=[nothing])).rows[0].warnings)
    assert "no MAC or IP" in w and "only one of" not in w


def test_parse_rejects_bad_files():
    with pytest.raises(CircuitImportError):
        parse_workbook(b"not an excel file")
    no_addr = [h if h not in ("Hardware Mac Address", "IP Address") else "Notes" for h in HEADER]
    with pytest.raises(CircuitImportError, match="Mac Address"):
        parse_workbook(make_xlsx(header=no_addr))
    wb = Workbook()
    wb.active.append(["a", "b"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(CircuitImportError, match="header"):
        parse_workbook(buf.getvalue())


@pytest.mark.parametrize("raw, norm", [
    ("Gi2/0/13", "gi2/0/13"), ("GigabitEthernet2/0/13", "gi2/0/13"), ("gi 2/0/13", "gi2/0/13"),
    ("Te1/1/1", "te1/1/1"), ("TenGigabitEthernet1/1/1", "te1/1/1"), ("Fa0/1", "fa0/1"),
    ("Eth1/10", "eth1/10"), ("Ethernet1/10", "eth1/10"), ("port1.0.1", "port1.0.1"),
    ("Po10", "po10"), ("Port-channel10", "po10"),
])
def test_normalize_port(raw, norm):
    assert normalize_port(raw) == norm


def test_normalize_mac():
    assert normalize_mac("02005E202269") == "02:00:5e:20:22:69"
    assert normalize_mac("02-00-5E-20-22-69") == normalize_mac("0200.5e20.2269")
    assert normalize_mac("12345") == ""


def test_switch_matching():
    assert switch_matches("ACC-SW01", "acc-sw01")
    assert switch_matches("acc-sw01.corp.local", "acc-sw01")
    assert switch_matches("SITEA-ACC-SW01", "acc-sw01")
    assert not switch_matches("ACC-SW011", "acc-sw01")
    names = {1: "SITEA-ACC-SW01", 2: "ACC-SW02", 3: "SITEB-ACC-SW02", 4: "core1"}
    m = match_switches(names, {"acc-sw01", "acc-sw02", "acc-sw09"})
    assert m["acc-sw01"] == ("matched", [1])
    assert m["acc-sw02"] == ("matched", [2])  # exact name wins over the suffix match
    assert m["acc-sw09"] == ("unmatched", [])
    assert match_switches({1: "A-SW1", 2: "B-SW1"}, {"sw1"})["sw1"][0] == "ambiguous"


# --- API ----------------------------------------------------------------------

def _upload(client, data, filename="circuits.xlsx"):
    return client.put("/api/circuits/imports/upload", params={"filename": filename},
                      content=data)


@pytest.fixture
def switch(admin, credential):
    r = admin.post("/api/devices", json={"name": "SITEA-ACC-SW01", "address": "10.20.0.11",
                                         "platform": "cisco_ios", "site": "SITEA",
                                         "credential_id": credential["id"]})
    assert r.status_code == 201, r.text
    return r.json()


def test_upload_preview_commit_and_query(admin, switch):
    empty = admin.get("/api/circuits").json()
    assert empty["import"] is None and empty["circuits"] == []
    r = _upload(admin, make_xlsx())
    assert r.status_code == 201, r.text
    imp = r.json()
    s = imp["summary"]
    assert imp["status"] == "pending" and imp["row_count"] == 5
    assert s["header_row"] == 3 and s["unmatched_switches"] == ["acc-sw02"]
    assert s["rows_unmatched"] == 1 and s["diff"]["added"] == 5
    assert admin.get("/api/circuits").json()["circuits"] == []  # nothing live until commit

    assert admin.post(f"/api/circuits/imports/{imp['id']}/commit").status_code == 200
    data = admin.get("/api/circuits").json()
    assert data["import"]["status"] == "active" and len(data["circuits"]) == 5
    first = data["circuits"][0]
    assert first["match"] == "matched" and first["switch_device_name"] == "SITEA-ACC-SW01"

    fd = admin.get(f"/api/circuits/for-device/{switch['id']}").json()
    assert fd["count"] == 4
    assert [(g["service"], g["vlan"], len(g["circuits"])) for g in fd["groups"]] == [
        ("ACS", "34", 2), ("BMS", "66", 1), ("Intercom", "36", 1)]

    csv = admin.get("/api/circuits/export.csv").text
    assert csv.splitlines()[0].startswith("IP Partition,Service,Engineering prefix")
    assert "DC-SITEA-GF-03" in csv and "SITEA-ACC-SW01" in csv


def test_reimport_shows_changes_and_keeps_history(admin, switch):
    first = _upload(admin, make_xlsx()).json()
    admin.post(f"/api/circuits/imports/{first['id']}/commit")

    rows = [list(r) for r in ROWS]
    rows[0][14] = "Gi2/0/20"          # moved port
    del rows[1]                       # removed
    rows.append(["SEC", "CCTV", None, "CAM-01", "Camera", "Axis", "ACCC8E000001", "Ground Floor",
                 "Lobby", "10.20.40.10", "255.255.255.0", "10.20.40.1", 40, "STATIC", "Gi1/0/5",
                 "ACC-SW01", None, None])  # added
    second = _upload(admin, make_xlsx(rows=rows)).json()
    d = second["summary"]["diff"]
    assert (d["added"], d["removed"], d["changed"]) == (1, 1, 1)
    assert d["changed_sample"][0]["fields"] == ["port"]
    assert d["changed_sample"][0]["before"]["port"] == "Gi2/0/13"
    assert second["summary"]["replaces_import"] == first["id"]

    admin.post(f"/api/circuits/imports/{second['id']}/commit")
    history = admin.get("/api/circuits/imports").json()
    assert [i["status"] for i in history] == ["active", "superseded"]
    assert len(admin.get("/api/circuits").json()["circuits"]) == 5


def test_discard_and_replace_pending(admin):
    a = _upload(admin, make_xlsx()).json()
    b = _upload(admin, make_xlsx()).json()  # a new upload replaces the pending preview
    assert admin.get(f"/api/circuits/imports/{a['id']}").status_code == 404
    assert admin.delete(f"/api/circuits/imports/{b['id']}").status_code == 204
    assert admin.get("/api/circuits/imports").json() == []
    c = _upload(admin, make_xlsx()).json()
    admin.post(f"/api/circuits/imports/{c['id']}/commit")
    assert admin.delete(f"/api/circuits/imports/{c['id']}").status_code == 409


def test_upload_errors(admin):
    r = _upload(admin, make_xlsx(), filename="list.csv")
    assert r.status_code == 422 and "Excel workbook" in r.json()["detail"]
    r = _upload(admin, b"a,b\n1,2\n", filename="list.csv")  # unknown headers: ask for the columns
    assert r.status_code == 200 and r.json()["needs_mapping"] is True
    assert _upload(admin, b"x", filename="list.txt").status_code == 422
    r = _upload(admin, b"garbage")
    assert r.status_code == 422 and "xlsx" in r.json()["detail"]


def test_viewer_can_read_but_not_import(client, app, admin, switch):
    imp = _upload(admin, make_xlsx()).json()
    admin.post(f"/api/circuits/imports/{imp['id']}/commit")
    from app.core.auth import VIEWER, User

    app.state.authenticator.authenticate = lambda u, p: User(username=u, role=VIEWER)
    client.post("/api/auth/logout")
    client.post("/api/auth/login", json={"username": "v", "password": "x"})
    assert len(client.get("/api/circuits").json()["circuits"]) == 5
    assert client.get(f"/api/circuits/for-device/{switch['id']}").status_code == 200
    assert _upload(client, make_xlsx()).status_code == 403
    assert client.post(f"/api/circuits/imports/{imp['id']}/commit").status_code == 403


# --- the admins' own columns, CSV uploads and blank templates -------------------------------

def test_custom_columns_are_read_shown_compared_and_exported(admin, switch):
    from tests.test_inventory_assets import _xlsx

    r = admin.post("/api/circuits/fields", json={"label": "Patch panel"})
    assert r.status_code == 201 and r.json()["key"] == "x_patch_panel"
    assert admin.post("/api/circuits/fields", json={"label": "patch PANEL"}).status_code == 422
    assert admin.post("/api/circuits/fields", json={"label": "Room / Area"}).status_code == 422
    assert admin.post("/api/circuits/fields", json={"label": "Live status"}).status_code == 422
    meta = admin.get("/api/circuits/meta").json()
    assert meta["can_manage"] and meta["fields"][-1] == {"key": "x_patch_panel", "label": "Patch panel",
                                                          "custom": True}

    header = ["Device Name", "Hardware Mac Address", "Patch panel"]
    first = _upload(admin, _xlsx([header, ["cam-1", "02005e000001", "PP1-01"]])).json()
    assert first["summary"]["columns"]["x_patch_panel"] == "Patch panel"
    assert "IP Address" in first["summary"]["missing_columns"]
    admin.post(f"/api/circuits/imports/{first['id']}/commit")
    assert admin.get("/api/circuits").json()["circuits"][0]["x_patch_panel"] == "PP1-01"

    second = _upload(admin, _xlsx([header, ["cam-1", "02005e000001", "PP1-02"]])).json()
    d = second["summary"]["diff"]
    assert d["changed"] == 1 and d["changed_sample"][0]["fields"] == ["x_patch_panel"]
    assert second["summary"]["labels"]["x_patch_panel"] == "Patch panel"
    admin.post(f"/api/circuits/imports/{second['id']}/commit")

    rows = admin.get("/api/circuits/export.csv").text.splitlines()
    assert "Patch panel" in rows[0].split(",") and "PP1-02" in rows[1]
    assert admin.put("/api/circuits/fields/x_patch_panel", json={"label": "Panel"}).status_code == 200
    assert admin.delete("/api/circuits/fields/x_patch_panel").status_code == 204
    assert "x_patch_panel" not in admin.get("/api/circuits").json()["circuits"][0]


def test_only_admins_add_columns(client, app, admin):
    from app.core.auth import VIEWER, User

    app.state.authenticator.authenticate = lambda u, p: User(username=u, role=VIEWER)
    client.post("/api/auth/logout")
    client.post("/api/auth/login", json={"username": "v", "password": "x"})
    assert client.post("/api/circuits/fields", json={"label": "Mine"}).status_code == 403
    assert client.get("/api/circuits/meta").json()["can_manage"] is False
    assert client.get("/api/circuits/template.csv").status_code == 200


def test_a_csv_file_is_read_like_a_workbook(admin, switch):
    text = ("Device Name;Hardware Mac Address;IP Address;VLAN;PORT;DEVICE\n"
            "cam-1;0200.5e00.0001;10.20.40.10;40;Gi1/0/5;ACC-SW01\n;;;;;\n")
    r = _upload(admin, text.encode("cp1252"), filename="list.csv")
    assert r.status_code == 201, r.text
    assert r.json()["sheet"] == "CSV" and r.json()["row_count"] == 1
    admin.post(f"/api/circuits/imports/{r.json()['id']}/commit")
    c = admin.get("/api/circuits").json()["circuits"][0]
    assert (c["vlan"], c["port"], c["match"]) == ("40", "Gi1/0/5", "matched")


def test_a_filled_in_template_imports(admin, switch):
    admin.post("/api/circuits/fields", json={"label": "Patch panel"})
    full = admin.get("/api/circuits/template.csv").text.lstrip("\ufeff").strip().split(",")
    assert full[:3] == ["IP Partition", "Service", "Engineering prefix"] and full[-1] == "Patch panel"
    # Only a few columns chosen: MAC and IP are always added, the order is kept.
    some = admin.get("/api/circuits/template.csv", params={"columns": "device_name,x_patch_panel,nope"})
    header = some.text.lstrip("\ufeff").strip().split(",")
    assert header == ["Hardware Mac Address", "IP Address", "Device Name", "Patch panel"]
    filled = some.text + "02005e000009,10.20.40.9,cam-9,PP2-07\n"
    r = _upload(admin, filled.encode(), filename="circuit-list-template.csv")
    assert r.status_code == 201, r.text
    admin.post(f"/api/circuits/imports/{r.json()['id']}/commit")
    assert admin.get("/api/circuits").json()["circuits"][0]["x_patch_panel"] == "PP2-07"

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(admin.get("/api/circuits/template.xlsx",
                                            params={"columns": "vlan"}).content))
    assert [c.value for c in wb["Circuits"][1]] == ["Hardware Mac Address", "IP Address", "VLAN"]
    assert wb.sheetnames == ["Circuits", "How to fill in"]
    assert admin.get("/api/circuits/template.pdf").status_code == 404
