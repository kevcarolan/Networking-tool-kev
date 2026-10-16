"""Found on the network: MACs on switch edge ports that aren't on the list, and adding
them to the list through the usual preview and commit."""

import io

from openpyxl import load_workbook

from app.tools.circuits import live
from tests.test_circuit_live import _scan
from tests.test_circuits import ROWS, _upload, make_xlsx, switch  # noqa: F401 - fixture


def _edges():
    from sqlalchemy import select

    from app.core.db import session_scope
    from app.tools.circuits.models import LiveMac

    with session_scope() as db:
        return set(db.scalars(select(LiveMac.mac_norm).where(LiveMac.uplink.is_(False))))


def test_found_lists_only_what_is_not_on_the_list(app, admin, switch, fw_collector):  # noqa: F811
    _scan(app, fw_collector)
    everything = {d["mac"] for d in admin.get("/api/circuits/discovered").json()["devices"]}
    assert everything == _edges()  # no list yet: every edge MAC, no uplinks

    imp = _upload(admin, make_xlsx(rows=ROWS[:2])).json()
    admin.post(f"/api/circuits/imports/{imp['id']}/commit")
    found = admin.get("/api/circuits/discovered").json()["devices"]
    macs = {d["mac"] for d in found}
    assert "02:00:5e:10:c0:2f" not in macs and "02:00:5e:10:c0:5e" not in macs
    assert macs == everything - {"02:00:5e:10:c0:2f", "02:00:5e:10:c0:5e"}
    d = found[0]
    assert d["switch"] == "SITEA-ACC-SW01" and d["port"] and d["device_id"] == switch["id"]


def test_adopt_makes_a_preview_then_the_list(app, admin, switch, fw_collector):  # noqa: F811
    imp = _upload(admin, make_xlsx(rows=ROWS[:1])).json()
    admin.post(f"/api/circuits/imports/{imp['id']}/commit")
    _scan(app, fw_collector)
    found = admin.get("/api/circuits/discovered").json()["devices"]
    pick = [found[0]["mac"], found[1]["mac"]]
    r = admin.post("/api/circuits/discovered/adopt",
                   json={"macs": pick, "service": "Printers", "device_type": "Printer"})
    assert r.status_code == 201, r.text
    pending = r.json()
    assert pending["status"] == "pending" and pending["row_count"] == 3
    assert pending["summary"]["diff"]["added"] == 2 and pending["summary"]["diff"]["removed"] == 0
    admin.post(f"/api/circuits/imports/{pending['id']}/commit")
    rows = admin.get("/api/circuits").json()["circuits"]
    added = [c for c in rows if c["service"] == "Printers"]
    assert {c["mac"] for c in added} == set(pick)
    assert all(c["live_status"] == live.MATCHES and c["device_type"] == "Printer" for c in added)
    left = {d["mac"] for d in admin.get("/api/circuits/discovered").json()["devices"]}
    assert not left & set(pick)
    assert admin.post("/api/circuits/discovered/adopt",
                      json={"macs": pick}).status_code == 422  # already on the list


def test_found_spreadsheet_uploads_as_a_list(app, admin, switch, fw_collector):  # noqa: F811
    _scan(app, fw_collector)
    data = admin.get("/api/circuits/discovered.xlsx").content
    ws = load_workbook(io.BytesIO(data)).active
    assert [c.value for c in ws[1]][:5] == ["Hardware Mac Address", "IP Address", "VLAN",
                                            "DEVICE", "PORT"]
    r = _upload(admin, data, filename="found.xlsx")
    assert r.status_code == 201 and r.json()["row_count"] == ws.max_row - 1


def test_found_needs_view_and_adopt_needs_manage(client, app, admin):
    from app.core.auth import VIEWER, User

    app.state.authenticator.authenticate = lambda u, p: User(username=u, role=VIEWER)
    client.post("/api/auth/logout")
    client.post("/api/auth/login", json={"username": "v", "password": "x"})
    assert client.get("/api/circuits/discovered").status_code == 200
    assert client.post("/api/circuits/discovered/adopt",
                       json={"macs": ["00:11:22:33:44:55"]}).status_code == 403
