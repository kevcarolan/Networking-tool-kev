"""Network inventory: every device appears by itself, anything else can be added, admins add
columns, a row can become a managed device, and spreadsheets go in and out."""

import csv
import io

import pytest
from openpyxl import Workbook, load_workbook

from app.core.db import session_scope
from app.tools.firmware_upgrade.models import DeviceFacts


def _rows(client):
    r = client.get("/api/inventory/assets")
    assert r.status_code == 200, r.text
    return r.json()


def _by_name(client, name):
    return next(r for r in _rows(client) if (r["values"]["name"] or "") == name)


def _add(client, **values):
    r = client.post("/api/inventory/assets", json={"values": values})
    assert r.status_code == 201, r.text
    return r.json()["id"]


def test_devices_appear_with_what_the_tools_know(admin, device):
    with session_scope() as db:
        db.add(DeviceFacts(device_id=device["id"], model="C9300-48P", serial="FOC123", version="17.9.4"))
    row = _by_name(admin, "core-sw1")
    assert row["managed"] and row["device"]["name"] == "core-sw1"
    assert row["auto"]["model"] == "C9300-48P" and row["auto"]["serial"] == "FOC123"
    assert row["auto"]["mgmt_ip"] == "10.0.0.1" and row["auto"]["manufacturer"] == "Cisco"
    assert row["values"]["type"] == "switch" and row["values"]["model"] == ""
    admin.put(f"/api/inventory/assets/{row['id']}", json={"values": {"model": "C9300-48P-A"}})
    assert _by_name(admin, "core-sw1")["values"]["model"] == "C9300-48P-A"  # typed wins
    assert len([r for r in _rows(admin) if r["managed"]]) == 1  # not added twice


def test_anything_can_be_recorded_and_changed(admin):
    asset = _add(admin, name="HQ-R01", type="rack", site="HQ", rack_units="42",
                 warranty_end="31/03/2027")
    row = _by_name(admin, "HQ-R01")
    assert not row["managed"] and row["values"]["rack_units"] == 42
    assert row["values"]["warranty_end"] == "2027-03-31" and row["values"]["status"] == "in service"
    r = admin.put(f"/api/inventory/assets/{asset}", json={"values": {"status": "Spare", "room": "Comms 1"}})
    assert r.json()["changed"] == ["Status", "Room"]
    assert admin.get("/api/audit").json()[0]["detail"] == "HQ-R01: Status, Room"
    assert admin.delete(f"/api/inventory/assets/{asset}").status_code == 204


@pytest.mark.parametrize("values, message", [
    ({"type": "rack"}, "Name: required"),
    ({"name": "x", "type": "spaceship"}, "Type: one of"),
    ({"name": "x", "rack_position": "top"}, "a whole number"),
    ({"name": "x", "purchase_date": "last year"}, "a date"),
    ({"name": "x", "colour": "red"}, "Unknown column"),
])
def test_values_are_checked(admin, values, message):
    r = admin.post("/api/inventory/assets", json={"values": values})
    assert r.status_code == 422 and message in r.json()["detail"]


def test_managed_rows_are_renamed_and_removed_on_config_backups(admin, device):
    row = _by_name(admin, "core-sw1")
    r = admin.put(f"/api/inventory/assets/{row['id']}", json={"values": {"name": "other"}})
    assert r.status_code == 422 and "Config backups" in r.json()["detail"]
    assert admin.delete(f"/api/inventory/assets/{row['id']}").status_code == 409
    admin.delete(f"/api/devices/{device['id']}")
    row = _by_name(admin, "core-sw1")
    assert not row["managed"]  # its details stay, as inventory only


def test_a_recorded_item_is_linked_when_the_device_is_added(admin):
    _add(admin, name="edge-sw9", rack="R5", mgmt_ip="10.0.0.9")
    admin.post("/api/devices", json={"name": "EDGE-SW9", "address": "10.0.0.9", "platform": "cisco_ios"})
    rows = [r for r in _rows(admin) if r["values"]["name"].lower() == "edge-sw9"]
    assert len(rows) == 1 and rows[0]["managed"] and rows[0]["values"]["rack"] == "R5"


def test_admins_add_columns_for_everyone(admin, app):
    from fastapi.testclient import TestClient

    from tests.test_users import _add as add_user
    from tests.test_users import _login

    r = admin.post("/api/inventory/fields", json={"label": "Power feed", "kind": "choice",
                                                  "choices": ["A", "B", "A+B"]})
    assert r.status_code == 201 and r.json()["key"] == "x_power_feed"
    assert admin.post("/api/inventory/fields", json={"label": "power FEED"}).status_code == 422
    assert admin.post("/api/inventory/fields", json={"label": "Serial number"}).status_code == 422
    assert admin.post("/api/inventory/fields", json={"label": "X", "kind": "choice"}).status_code == 422
    asset = _add(admin, name="ups1", x_power_feed="a+b")
    assert _by_name(admin, "ups1")["values"]["x_power_feed"] == "A+B"
    assert admin.post("/api/inventory/assets", json={"values": {"name": "u2", "x_power_feed": "C"}}).status_code == 422
    admin.put("/api/inventory/fields/x_power_feed", json={"label": "Feed", "kind": "choice", "choices": ["A", "B", "A+B"]})
    meta = admin.get("/api/inventory/meta").json()
    assert {"key": "x_power_feed", "label": "Feed"}.items() <= next(
        f for f in meta["fields"] if f["key"] == "x_power_feed").items()

    add_user(admin, "editor", permissions={"inventory": "manage"})
    with TestClient(app) as c:
        _login(c, "editor")
        assert c.post("/api/inventory/fields", json={"label": "Mine"}).status_code == 403
        assert c.put(f"/api/inventory/assets/{asset}", json={"values": {"x_power_feed": "B"}}).status_code == 200

    assert admin.delete("/api/inventory/fields/x_power_feed").status_code == 204
    assert "x_power_feed" not in _by_name(admin, "ups1")["values"]


def test_viewers_see_but_dont_change(app, admin):
    from fastapi.testclient import TestClient

    from tests.test_users import _add as add_user
    from tests.test_users import _login

    _add(admin, name="HQ-R01", type="rack")
    add_user(admin, "viewer1", preset="viewer")
    with TestClient(app) as c:
        _login(c, "viewer1")
        assert len(_rows(c)) == 1 and c.get("/api/inventory/meta").json()["can_manage"] is False
        assert c.post("/api/inventory/assets", json={"values": {"name": "x"}}).status_code == 403


def test_a_recorded_item_becomes_a_managed_device(admin, credential):
    asset = _add(admin, name="new-sw", site="HQ", mgmt_ip="10.0.0.20")
    r = admin.post(f"/api/inventory/assets/{asset}/manage", json={
        "platform": "cisco_ios", "address": "10.0.0.20", "credential_id": credential["id"]})
    assert r.status_code == 200, r.text
    devices = admin.get("/api/devices").json()
    assert [d["name"] for d in devices] == ["new-sw"] and devices[0]["site"] == "HQ"
    assert _by_name(admin, "new-sw")["managed"]
    assert admin.post(f"/api/inventory/assets/{asset}/manage", json={
        "platform": "cisco_ios", "address": "10.0.0.20"}).status_code == 409


def test_managing_needs_the_addresses_level_and_a_free_name(app, admin, device):
    from fastapi.testclient import TestClient

    from tests.test_users import _add as add_user
    from tests.test_users import _login

    asset = _add(admin, name="core-sw1-spare")
    clash = _add(admin, name="CORE-SW1 ")
    assert admin.post(f"/api/inventory/assets/{asset}/manage",
                      json={"platform": "nope", "address": "x"}).status_code == 422
    add_user(admin, "details", permissions={"inventory": "manage"})
    with TestClient(app) as c:
        _login(c, "details")
        assert c.post(f"/api/inventory/assets/{asset}/manage",
                      json={"platform": "cisco_ios", "address": "10.0.0.30"}).status_code == 403
    r = admin.post(f"/api/inventory/assets/{clash}/manage", json={"platform": "cisco_ios", "address": "10.0.0.31"})
    assert r.status_code == 409 and "already exists" in r.json()["detail"]


def _xlsx(rows):
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _import(client, data, name, commit=False):
    return client.put("/api/inventory/import", params={"filename": name, "commit": commit}, content=data)


def test_import_shows_what_changes_then_makes_the_changes(admin, device):
    _add(admin, name="HQ-UPS1", serial="AS123", type="UPS")
    data = _xlsx([["Name", "Type", "Serial number", "Rack", "Warranty end", "Colour"],
                  ["hq-ups-renamed", "UPS", "AS123", "HQ-R01", "2027-03-31", "grey"],
                  ["core-sw1", "", "", "HQ-R02", "", ""],
                  ["HQ-PDU-A", "PDU", "", "HQ-R01", "", ""],
                  ["", "", "", "", "", ""]])
    plan = _import(admin, data, "inv.xlsx").json()
    assert [n["name"] for n in plan["new"]] == ["HQ-PDU-A"]
    assert {u["name"]: u["changes"] for u in plan["updated"]} == {
        "HQ-UPS1": ["Name", "Rack", "Warranty end"], "core-sw1": ["Rack"]}
    assert plan["unknown_columns"] == ["Colour"] and plan["errors"] == []
    assert len(_rows(admin)) == 2  # nothing changed yet
    done = _import(admin, data, "inv.xlsx", commit=True).json()
    assert done["committed"]
    assert _by_name(admin, "hq-ups-renamed")["values"]["rack"] == "HQ-R01"
    assert _by_name(admin, "core-sw1")["values"]["rack"] == "HQ-R02"  # matched by name, still managed
    assert admin.get("/api/audit").json()[0]["action"] == "inventory.import"


def test_a_file_with_errors_isnt_imported(admin):
    data = "Name,Type,Position (U)\nok,rack,1\nbad,rack,top\nok,rack,2\n".encode()
    plan = _import(admin, data, "inv.csv").json()
    assert [e["row"] for e in plan["errors"]] == [3, 4]
    assert _import(admin, data, "inv.csv", commit=True).status_code == 422
    assert _rows(admin) == []
    assert _import(admin, b"Type\nrack\n", "x.csv").status_code == 422  # no Name column
    assert _import(admin, b"PK\x03\x04junk", "x.xlsx").status_code == 422
    assert _import(admin, b"a", "x.txt").status_code == 422


def test_exports_have_every_column_and_are_safe_in_excel(admin, device):
    admin.post("/api/inventory/fields", json={"label": "Power feed"})
    _add(admin, name="=HYPERLINK(\"http://x\")", x_power_feed="A")
    rows = list(csv.reader(io.StringIO(admin.get("/api/inventory/export.csv").text)))
    assert rows[0][:3] == ["Name", "Type", "Role"] and rows[0][-2:] == ["Power feed", "Managed"]
    names = [r[0] for r in rows[1:]]
    assert "core-sw1" in names and "'=HYPERLINK(\"http://x\")" in names
    managed = next(r for r in rows[1:] if r[0] == "core-sw1")
    assert managed[-1] == "yes" and "10.0.0.1" in managed  # the address, from the device
    ws = load_workbook(io.BytesIO(admin.get("/api/inventory/export.xlsx").content)).active
    assert ws["A1"].value == "Name" and any(c.value == "Power feed" for c in ws[1])
    assert all(c.data_type == "s" for c in ws["A"] if str(c.value).startswith("="))


def test_an_export_imports_back_unchanged(admin, device):
    _add(admin, name="HQ-R01", type="rack", rack_units=42)
    data = admin.get("/api/inventory/export.xlsx").content
    plan = _import(admin, data, "inventory.xlsx").json()
    assert plan["new"] == [] and plan["errors"] == []


def test_a_blank_template_has_the_chosen_columns_and_imports(admin):
    admin.post("/api/inventory/fields", json={"label": "Power feed", "kind": "choice",
                                              "choices": ["A", "B"]})
    full = admin.get("/api/inventory/template.csv").text.lstrip("\ufeff").strip().split(",")
    assert full[:2] == ["Name", "Type"] and full[-1] == "Power feed" and "Managed" not in full
    some = admin.get("/api/inventory/template.csv", params={"columns": "rack,x_power_feed"}).text
    assert some.lstrip("\ufeff").strip() == "Name,Rack,Power feed"  # Name is always there
    plan = _import(admin, (some + "HQ-UPS1,HQ-R01,B\n").encode(), "t.csv", commit=True).json()
    assert plan["committed"] and _by_name(admin, "HQ-UPS1")["values"]["x_power_feed"] == "B"

    wb = load_workbook(io.BytesIO(admin.get("/api/inventory/template.xlsx",
                                            params={"columns": "type,x_power_feed"}).content))
    ws = wb["Inventory"]
    assert [c.value for c in ws[1]] == ["Name", "Type", "Power feed"]
    lists = sorted(dv.formula1 for dv in ws.data_validations.dataValidation)
    assert lists[0] == '"A,B"' and "rack" in lists[1]
    assert _import(admin, admin.get("/api/inventory/template.xlsx").content, "t.xlsx").json()["new"] == []
