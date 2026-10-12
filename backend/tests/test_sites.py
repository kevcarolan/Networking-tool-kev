"""Sites: the list of sites, what the inventory has at each, columns, and spreadsheets."""

import io

from openpyxl import load_workbook

from app.core.db import session_scope
from app.tools.config_backup.models import BackupState
from tests.test_inventory_assets import _add as add_item
from tests.test_inventory_assets import _xlsx


def _add(client, **values):
    r = client.post("/api/sites", json={"values": values})
    assert r.status_code == 201, r.text
    return r.json()["id"]


def _site(client, name):
    return next(s for s in client.get("/api/sites").json()["sites"] if s["values"]["name"] == name)


def test_sites_count_what_the_inventory_has_there(admin, device):
    # device: core-sw1 at site "HQ" (from the device itself)
    with session_scope() as db:
        db.add(BackupState(device_id=device["id"], frequency_minutes=1440, last_status="failed"))
    _add(admin, name="Head office", code="HQ", type="Data centre", city="Leeds")
    add_item(admin, name="HQ-R01", type="rack", site="head OFFICE")  # by name, any case
    add_item(admin, name="DC2-UPS1", type="UPS", site="DC2")         # not on the list
    data = admin.get("/api/sites").json()
    lab = data["sites"][0]
    assert lab["values"]["type"] == "data centre" and lab["values"]["status"] == "live"
    assert lab["counts"] == {"items": 2, "devices": 1, "backup_failed": 1, "behind": 0, "circuits": 0}
    assert data["unlisted"] == [{"name": "DC2", "items": 1}]
    items = admin.get(f"/api/sites/{lab['id']}/items").json()
    assert sorted(i["name"] for i in items) == ["HQ-R01", "core-sw1"]

    assert admin.post("/api/sites/from-inventory").json()["added"] == ["DC2"]
    assert admin.get("/api/sites").json()["unlisted"] == []
    assert admin.get("/api/inventory/meta").json()["sites"] == ["DC2", "Head office"]


def test_names_and_codes_are_unique_and_checked(admin):
    _add(admin, name="Data Centre 1", code="DC1")
    assert admin.post("/api/sites", json={"values": {"name": "dc1"}}).status_code == 409
    assert admin.post("/api/sites", json={"values": {"name": "B", "code": "data centre 1"}}).status_code == 409
    assert admin.post("/api/sites", json={"values": {"code": "X"}}).status_code == 422
    r = admin.post("/api/sites", json={"values": {"name": "B", "type": "moon base"}})
    assert r.status_code == 422 and "Type: one of" in r.json()["detail"]


def test_renaming_a_site_moves_its_items_and_devices(admin, device):
    site = _add(admin, name="Head office", code="HQ")
    add_item(admin, name="rack-1", site="head office")
    r = admin.put(f"/api/sites/{site}", json={"values": {"name": "Leeds Lab", "code": "LDS"}})
    assert r.json() == {"changed": ["Name", "Site code"], "renamed_on": 2}
    assert admin.get("/api/devices").json()[0]["site"] == "LDS"  # it used the code
    assert _site(admin, "Leeds Lab")["counts"]["items"] == 2
    assert "Site changed on 2" in admin.get("/api/audit").json()[0]["detail"]
    assert admin.delete(f"/api/sites/{site}").status_code == 204
    assert admin.get("/api/devices").json()[0]["site"] == "LDS"  # items keep their value


def test_viewers_see_sites_but_dont_change_them(app, admin):
    from fastapi.testclient import TestClient

    from tests.test_users import _add as add_user
    from tests.test_users import _login

    _add(admin, name="HQ")
    add_user(admin, "viewer1", preset="viewer")
    with TestClient(app) as c:
        _login(c, "viewer1")
        assert len(c.get("/api/sites").json()["sites"]) == 1
        assert c.get("/api/sites/meta").json()["can_manage"] is False
        assert c.post("/api/sites", json={"values": {"name": "x"}}).status_code == 403
        assert c.post("/api/sites/fields", json={"label": "x"}).status_code == 403


def test_custom_columns_templates_import_and_export(admin):
    r = admin.post("/api/sites/fields", json={"label": "Power", "kind": "choice", "choices": ["A", "A+B"]})
    assert r.status_code == 201 and r.json()["key"] == "x_power"
    assert admin.post("/api/sites/fields", json={"label": "Site code"}).status_code == 422
    _add(admin, name="Data Centre 1", code="DC1")

    head = admin.get("/api/sites/template.csv", params={"columns": "code,x_power"}).text
    assert head.lstrip("﻿").strip() == "Name,Site code,Power"
    data = (head + "DC 1 (renamed),DC1,A+B\nOffice 2,OF2,a\nbad,,C\n").encode()
    plan = admin.put("/api/sites/import", params={"filename": "s.csv"}, content=data).json()
    assert [n["name"] for n in plan["new"]] == ["Office 2"]
    assert plan["updated"][0]["changes"] == ["Name", "Power"]  # matched by code
    assert [e["row"] for e in plan["errors"]] == [4]
    assert admin.put("/api/sites/import", params={"filename": "s.csv", "commit": True},
                     content=data).status_code == 422
    good = data.replace(b"bad,,C\n", b"")
    done = admin.put("/api/sites/import", params={"filename": "s.csv", "commit": True}, content=good).json()
    assert done["committed"] and _site(admin, "Office 2")["values"]["x_power"] == "A"

    rows = admin.get("/api/sites/export.csv").text.splitlines()
    assert rows[0].startswith("Name,Site code,Type") and rows[0].endswith("Backups failing,Firmware behind,Circuits")
    ws = load_workbook(io.BytesIO(admin.get("/api/sites/export.xlsx").content)).active
    assert ws["A2"].value == "DC 1 (renamed)"
    again = admin.put("/api/sites/import", params={"filename": "sites.xlsx"},
                      content=admin.get("/api/sites/export.xlsx").content).json()
    assert again["new"] == [] and again["updated"] == [] and again["unknown_columns"] == []
    xlsx = _xlsx([["Name", "Type"], ["Depot", "warehouse"]])
    assert admin.put("/api/sites/import", params={"filename": "x.xlsx"}, content=xlsx).json()["new"][0]["name"] == "Depot"
    assert admin.delete("/api/sites/fields/x_power").status_code == 204


def test_circuits_are_linked_to_sites(admin, device):
    """The list's Site column first, else the site of the switch the circuit is on."""
    from tests.test_circuits import _upload

    hq = _add(admin, name="Head office", code="HQ")
    dc = _add(admin, name="Data centre 1", code="DC1")
    data = _xlsx([["Device Name", "Hardware Mac Address", "Site", "DEVICE", "PORT"],
                  ["cam-1", "02005e000001", "dc1", "", ""],          # the list says DC1
                  ["cam-2", "02005e000002", "", "core-sw1", "Gi1/0/2"],  # on core-sw1, at HQ
                  ["cam-3", "02005e000003", "Moon base", "", ""],    # not on the site list
                  ["cam-4", "02005e000004", "", "", ""]])            # no site at all
    imp = _upload(admin, data).json()
    admin.post(f"/api/circuits/imports/{imp['id']}/commit")
    rows = {c["device_name"]: c for c in admin.get("/api/circuits").json()["circuits"]}
    assert (rows["cam-1"]["site_id"], rows["cam-1"]["site_from"]) == (dc, "list")
    assert (rows["cam-2"]["site_name"], rows["cam-2"]["site_from"]) == ("Head office", "switch")
    assert (rows["cam-3"]["site_id"], rows["cam-3"]["site_text"]) == (None, "Moon base")
    assert rows["cam-4"]["site_from"] == ""
    counts = {s["values"]["name"]: s["counts"]["circuits"] for s in admin.get("/api/sites").json()["sites"]}
    assert counts == {"Data centre 1": 1, "Head office": 1}
    assert "Circuits" in [c["label"] for c in admin.get("/api/sites/meta").json()["counts"]]
    assert admin.get("/api/sites/index").json() == [
        {"id": dc, "name": "Data centre 1", "code": "DC1"}, {"id": hq, "name": "Head office", "code": "HQ"}]

    # Renaming the code carries to the circuit list, so the circuit stays at its site.
    admin.put(f"/api/sites/{dc}", json={"values": {"code": "DC-1"}})
    rows = {c["device_name"]: c for c in admin.get("/api/circuits").json()["circuits"]}
    assert rows["cam-1"]["site"] == "DC-1" and rows["cam-1"]["site_id"] == dc
    csv_head = admin.get("/api/circuits/export.csv").text.splitlines()[0].split(",")
    assert "Site" in csv_head and "Linked site" in csv_head
